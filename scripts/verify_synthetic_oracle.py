"""Independent W06 oracle for the harmless selected-report fixtures.

This utility deliberately imports no Metroliza modules.  It compares an eventual
fresh acceptance SQLite database, a produced XLSX workbook and a captured grouping
snapshot against values manually transcribed from the authored PDF fixture and its
selected filenames.  ``--self-check`` creates only synthetic stand-ins to validate
this comparator; it is not an application, package, or OCR execution.
"""
from __future__ import annotations

import argparse
import json
import shutil
import math
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
except ImportError as exc:  # pragma: no cover - environment prerequisite
    raise SystemExit("openpyxl is required for the standalone XLSX comparator") from exc


@dataclass(frozen=True)
class OracleMismatch(AssertionError):
    scope: str
    detail: str

    def __str__(self) -> str:
        return f"{self.scope}: {self.detail}"


def _fail(scope: str, detail: str) -> None:
    raise OracleMismatch(scope, detail)


def _number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        _fail("value", "non-numeric value")


def _same(expected: Any, actual: Any, *, scope: str) -> None:
    if isinstance(expected, float):
        parsed = _number(actual)
        if parsed is None or not math.isfinite(parsed) or abs(expected - parsed) > 1e-9:
            _fail(scope, "numeric value differs")
    elif expected is None:
        if actual is not None:
            _fail(scope, "expected empty value")
    elif isinstance(expected, str) and (not isinstance(actual, str) or expected != actual):
        _fail(scope, "text value differs")
    elif expected != actual:
        _fail(scope, "text value differs")


def _load_oracle(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if data.get("schema") != "metroliza-synthetic-report-oracle-v2":
        _fail("oracle", "unsupported schema")
    return data


def _database_rows(connection: sqlite3.Connection) -> list[dict[str, Any]]:
    query = """
        SELECT pr.id AS report_id, m.id AS measurement_id, sfl.file_name, sf.sha256,
               pr.measurement_count, pr.has_nok, rm.reference, rm.report_date, rm.sample_number,
               rm.sample_number_kind, m.row_order, m.header, m.section_name,
               m.feature_label, m.characteristic_name, m.characteristic_family,
               m.description, m.ax, m.nominal, m.tol_plus, m.tol_minus, m.bonus,
               m.meas, m.dev, m.outtol, m.is_nok, m.status_code
        FROM parsed_reports pr
        JOIN source_files sf ON sf.id = pr.source_file_id
        JOIN source_file_locations sfl ON sfl.source_file_id = sf.id AND sfl.is_active = 1
        JOIN report_metadata rm ON rm.report_id = pr.id
        JOIN report_measurements m ON m.report_id = pr.id
        ORDER BY pr.id, m.id
    """
    columns = [item[0] for item in connection.execute(query).description]
    return [dict(zip(columns, row, strict=True)) for row in connection.execute(query)]


def assert_database(oracle: dict[str, Any], database: Path) -> None:
    with sqlite3.connect(database.resolve().as_uri() + "?mode=ro", uri=True) as connection:
        columns = {row[1] for row in connection.execute("PRAGMA table_info(report_measurements)")}
        if "unit" in columns:
            _fail("database.unit", "fixture has no unit and schema must not invent a unit field")
        expected_ids = [row["report_id"] for row in oracle["persisted_reports"]]
        table_queries = {
            "source_files": "SELECT id FROM source_files ORDER BY id",
            "active_locations": "SELECT source_file_id FROM source_file_locations WHERE is_active = 1 ORDER BY source_file_id",
            "parsed_reports": "SELECT id FROM parsed_reports ORDER BY id",
            "metadata": "SELECT report_id FROM report_metadata ORDER BY report_id",
            "measurements": "SELECT report_id FROM report_measurements ORDER BY report_id, id",
        }
        for name, query in table_queries.items():
            actual_ids = [row[0] for row in connection.execute(query)]
            required = expected_ids if name != "measurements" else expected_ids
            if actual_ids != required:
                _fail(f"database.{name}", "unexpected keyset or cardinality")
        locations = connection.execute(
            "SELECT source_file_id, is_active FROM source_file_locations ORDER BY source_file_id"
        ).fetchall()
        if locations != [(identifier, 1) for identifier in expected_ids]:
            _fail("database.locations", "unexpected inactive, duplicate or orphan location")
        actual = _database_rows(connection)

    expected_reports = oracle["persisted_reports"]
    expected_inputs = oracle["fixture_construction"]["staged_aliases"]
    if len(actual) != len(expected_reports):
        _fail("database.rows", "unexpected persisted measurement count")

    for index, (report, source, row) in enumerate(zip(expected_reports, expected_inputs, actual, strict=True)):
        fields = {
            "report_id": report["report_id"],
            "measurement_id": report["measurement"]["measurement_id"],
            "file_name": source["file_name"],
            "sha256": source["sha256"],
            "measurement_count": report["measurement_count"],
            "has_nok": report["has_nok"],
            "reference": report["reference"],
            "report_date": report["report_date"],
            "sample_number": report["sample_number"],
            "sample_number_kind": report["sample_number_kind"],
        }
        fields.update(report["measurement"])
        for field, expected in fields.items():
            _same(expected, row[field], scope=f"database.row{index + 1}.{field}")


def _find_measurement_blocks(workbook: Path, headers: list[str]) -> list[tuple[Any, int, int]]:
    book = load_workbook(workbook, data_only=False, read_only=True)
    found = []
    try:
        for sheet in book.worksheets:
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                for start in range(len(row) - len(headers) + 1):
                    if [str(value) if value is not None else "" for value in row[start:start + len(headers)]] == headers:
                        found.append((sheet.title, row_index, start + 1))
    finally:
        book.close()
    return found


def assert_workbook(oracle: dict[str, Any], workbook: Path) -> None:
    expected = oracle["workbook_output"]["measurement_block"]
    headers = expected["headers"]
    matches = _find_measurement_blocks(workbook, headers)
    if len(matches) != 1:
        _fail("workbook.block", "measurement header block must occur exactly once")
    sheet_name, header_row, start_column = matches[0]
    if header_row != expected["data_header_excel_row"]:
        _fail("workbook.block", "measurement block is on the wrong row")

    book = load_workbook(workbook, data_only=False, read_only=True)
    try:
        sheet = book[sheet_name]
        for row_offset, expected_row in enumerate(expected["rows"], start=1):
            for column_offset, expected_value in enumerate(expected_row):
                _same(
                    expected_value,
                    sheet.cell(row=header_row + row_offset, column=start_column + column_offset).value,
                    scope=f"workbook.data.row{row_offset}.column{column_offset}",
                )
        if any(
            sheet.cell(row=header_row + len(expected["rows"]) + 1, column=start_column + offset).value is not None
            for offset in range(len(headers))
        ):
            _fail("workbook.data", "unexpected third measurement row")
        for row_number, expected_row in enumerate(expected["static_limits"], start=1):
            _same(expected_row[0], sheet.cell(row=row_number, column=start_column).value, scope=f"workbook.static.row{row_number}.label")
            _same(expected_row[1], sheet.cell(row=row_number, column=start_column + 1).value, scope=f"workbook.static.row{row_number}.value")
    finally:
        book.close()


def assert_grouping(oracle: dict[str, Any], grouping: Path) -> None:
    expected = oracle["grouping_output"]
    actual = json.loads(grouping.read_text(encoding="utf-8"))
    if actual != expected:
        _fail("grouping", "semantic group membership or output metric differs")


def assert_tabular(oracle: dict[str, Any], tabular: Path) -> None:
    actual = json.loads(tabular.read_text(encoding="utf-8"))
    # Canonical JSON also distinguishes booleans/floats from integer row IDs.
    if json.dumps(actual, sort_keys=True) != json.dumps(oracle["tabular_output"], sort_keys=True):
        _fail("tabular", "filter IDs, source lexemes/types or fixture identity differ")


def verify(oracle_path: Path, database: Path, workbook: Path, grouping: Path,
           tabular: Path | None = None) -> None:
    oracle = _load_oracle(oracle_path)
    assert_database(oracle, database)
    assert_workbook(oracle, workbook)
    assert_grouping(oracle, grouping)
    if tabular is not None:
        assert_tabular(oracle, tabular)


def _create_synthetic_database(oracle: dict[str, Any], path: Path) -> None:
    if path.exists():
        path.unlink()
    with sqlite3.connect(path) as connection:
        connection.executescript(
            """
            CREATE TABLE source_files (id INTEGER PRIMARY KEY, sha256 TEXT NOT NULL);
            CREATE TABLE source_file_locations (source_file_id INTEGER, file_name TEXT, is_active INTEGER);
            CREATE TABLE parsed_reports (id INTEGER PRIMARY KEY, source_file_id INTEGER, measurement_count INTEGER, has_nok INTEGER);
            CREATE TABLE report_metadata (report_id INTEGER PRIMARY KEY, reference TEXT, report_date TEXT, sample_number TEXT, sample_number_kind TEXT);
            CREATE TABLE report_measurements (
                id INTEGER PRIMARY KEY, report_id INTEGER, row_order INTEGER, header TEXT, section_name TEXT,
                feature_label TEXT, characteristic_name TEXT, characteristic_family TEXT, description TEXT,
                ax TEXT, nominal REAL, tol_plus REAL, tol_minus REAL, bonus REAL, meas REAL, dev REAL,
                outtol REAL, is_nok INTEGER, status_code TEXT
            );
            """
        )
        for report, source in zip(oracle["persisted_reports"], oracle["fixture_construction"]["staged_aliases"], strict=True):
            measurement = report["measurement"]
            report_id = report["report_id"]
            connection.execute("INSERT INTO source_files VALUES (?, ?)", (report_id, source["sha256"]))
            connection.execute("INSERT INTO source_file_locations VALUES (?, ?, 1)", (report_id, source["file_name"]))
            connection.execute("INSERT INTO parsed_reports VALUES (?, ?, ?, ?)", (report_id, report_id, report["measurement_count"], report["has_nok"]))
            connection.execute("INSERT INTO report_metadata VALUES (?, ?, ?, ?, ?)", (report_id, report["reference"], report["report_date"], report["sample_number"], report["sample_number_kind"]))
            values = [measurement[key] for key in (
                "measurement_id", "row_order", "header", "section_name", "feature_label", "characteristic_name",
                "characteristic_family", "description", "ax", "nominal", "tol_plus", "tol_minus", "bonus",
                "meas", "dev", "outtol", "is_nok", "status_code",
            )]
            measurement_values = [measurement["measurement_id"], report_id, *values[1:]]
            connection.execute(
                "INSERT INTO report_measurements VALUES (" + ", ".join("?" for _ in measurement_values) + ")",
                measurement_values,
            )


def _create_synthetic_workbook(oracle: dict[str, Any], path: Path) -> None:
    block = oracle["workbook_output"]["measurement_block"]
    book = Workbook()
    sheet = book.active
    sheet.title = "Results"
    for row_number, (label, value) in enumerate(block["static_limits"], start=1):
        sheet.cell(row=row_number, column=1, value=label)
        sheet.cell(row=row_number, column=2, value=value)
    for column, header in enumerate(block["headers"], start=1):
        sheet.cell(row=block["data_header_excel_row"], column=column, value=header)
    for offset, row in enumerate(block["rows"], start=1):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=block["data_header_excel_row"] + offset, column=column, value=value)
    book.save(path)


def _expect_mismatch(label: str, action) -> str:
    try:
        action()
    except OracleMismatch as error:
        return error.scope
    raise AssertionError(f"negative control {label} unexpectedly passed")


def self_check(oracle_path: Path, output: Path) -> dict[str, Any]:
    output.mkdir(parents=True, exist_ok=True)
    oracle = _load_oracle(oracle_path)
    database = output / "oracle-ok.sqlite"
    workbook = output / "oracle-ok.xlsx"
    grouping = output / "oracle-ok-grouping.json"
    _create_synthetic_database(oracle, database)
    _create_synthetic_workbook(oracle, workbook)
    grouping.write_text(json.dumps(oracle["grouping_output"], indent=2) + "\n", encoding="utf-8")
    verify(oracle_path, database, workbook, grouping)

    bad_database = output / "oracle-corrupt-measurement.sqlite"
    shutil.copyfile(database, bad_database)
    with sqlite3.connect(bad_database) as connection:
        connection.execute("UPDATE report_measurements SET meas = 10.03 WHERE id = 2")
    db_scope = _expect_mismatch("database", lambda: verify(oracle_path, bad_database, workbook, grouping))

    orphan_database = output / "oracle-orphan.sqlite"
    shutil.copyfile(database, orphan_database)
    with sqlite3.connect(orphan_database) as connection:
        connection.execute("INSERT INTO source_files VALUES (?, ?)", (9, "orphan"))
    orphan_scope = _expect_mismatch("orphan", lambda: verify(oracle_path, orphan_database, workbook, grouping))

    nan_database = output / "oracle-nan.sqlite"
    shutil.copyfile(database, nan_database)
    with sqlite3.connect(nan_database) as connection:
        connection.execute("UPDATE report_measurements SET meas = ? WHERE id = 2", (float("nan"),))
    nan_scope = _expect_mismatch("nan", lambda: verify(oracle_path, nan_database, workbook, grouping))

    bad_workbook = output / "oracle-corrupt-output.xlsx"
    shutil.copyfile(workbook, bad_workbook)
    book = load_workbook(bad_workbook)
    book["Results"].cell(row=23, column=3, value=10.03)
    book.save(bad_workbook)
    workbook_scope = _expect_mismatch("workbook", lambda: verify(oracle_path, database, bad_workbook, grouping))

    duplicate_workbook = output / "oracle-duplicate-block.xlsx"
    shutil.copyfile(workbook, duplicate_workbook)
    book = load_workbook(duplicate_workbook)
    for column, value in enumerate(oracle["workbook_output"]["measurement_block"]["headers"], start=1):
        book["Results"].cell(row=30, column=column, value=value)
    book.save(duplicate_workbook)
    duplicate_scope = _expect_mismatch("duplicate", lambda: verify(oracle_path, database, duplicate_workbook, grouping))

    bad_grouping = output / "oracle-corrupt-grouping.json"
    corrupt_grouping = json.loads(grouping.read_text(encoding="utf-8"))
    corrupt_grouping["projection"]["members"]["POPULATION"]["report_ids"] = [1, 9]
    bad_grouping.write_text(json.dumps(corrupt_grouping, indent=2) + "\n", encoding="utf-8")
    grouping_scope = _expect_mismatch("grouping", lambda: verify(oracle_path, database, workbook, bad_grouping))

    return {
        "schema": "metroliza-synthetic-oracle-self-check-v1",
        "result": "passed",
        "kind": "comparator_self_validation_not_package_acceptance",
        "positive_control": "passed",
        "negative_controls": {
            "database": db_scope, "orphan_database": orphan_scope, "nan_database": nan_scope,
            "workbook": workbook_scope, "duplicate_workbook_block": duplicate_scope, "grouping": grouping_scope,
        },
        "limits": ["No Metroliza module, parser, GUI, package, OCR, or real data was executed."],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--oracle", type=Path, required=True)
    parser.add_argument("--database", type=Path)
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--grouping", type=Path)
    parser.add_argument("--tabular", type=Path)
    parser.add_argument("--self-check", type=Path)
    args = parser.parse_args()
    if args.self_check is not None:
        if any(value is not None for value in (args.database, args.workbook, args.grouping, args.tabular)):
            raise SystemExit("--self-check cannot be combined with acceptance artifacts")
        print(json.dumps(self_check(args.oracle, args.self_check), sort_keys=True))
        return
    if any(value is None for value in (args.database, args.workbook, args.grouping, args.tabular)):
        raise SystemExit("--database, --workbook, --grouping and --tabular are required for an acceptance comparison")
    verify(args.oracle, args.database, args.workbook, args.grouping, args.tabular)
    print(json.dumps({"result": "passed", "kind": "acceptance_artifact_comparison"}, sort_keys=True))


if __name__ == "__main__":
    main()
