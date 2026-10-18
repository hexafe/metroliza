"""Equivalence tests for the bulk worksheet path (#1124).

Every case writes the same rows twice -- cell by cell through
``write_untrusted_xlsx_cell`` and through ``BulkRowsWorksheet.write_bulk_block``
-- and requires byte-identical workbook parts.
"""

import math
import random
import tempfile
import types
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import xlsxwriter
import numpy as np

from metroliza.exporting import xlsx_bulk_sheet
from metroliza.exporting.export_backends import _is_blank_cell
from metroliza.exporting.xlsx_bulk_sheet import (
    BULK_WRITE_ENV,
    BulkRowsWorksheet,
    bulk_write_enabled,
    enable_bulk_rows,
    plan_bulk_rows,
)
from metroliza.exporting.xlsx_writer_policy import write_untrusted_xlsx_cell, xlsxwriter_workbook_options

_VOLATILE = {'docProps/core.xml'}


def _read_parts(path):
    with zipfile.ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist() if name not in _VOLATILE}


def _write_classic(worksheet, columns, rows):
    for column_index, name in enumerate(columns):
        write_untrusted_xlsx_cell(worksheet, 0, column_index, name)
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row):
            write_untrusted_xlsx_cell(worksheet, row_index, column_index, None if _is_blank_cell(value) else value)


def _write_bulk(worksheet, columns, rows):
    plan = plan_bulk_rows(columns, rows)
    if plan is None:
        raise AssertionError('rows were expected to be bulk-writable')
    for column_index, name in enumerate(columns):
        write_untrusted_xlsx_cell(worksheet, 0, column_index, name)
    if not worksheet.write_bulk_block(1, plan):
        raise AssertionError('bulk block was refused')


class BulkEquivalenceCase(unittest.TestCase):
    def build(self, path, *, bulk, columns, rows, extra=None, prelude=None):
        book = xlsxwriter.Workbook(str(path), xlsxwriter_workbook_options())
        if bulk:
            self.assertTrue(enable_bulk_rows(book))
        if prelude is not None:
            prelude(book)
        sheet = book.add_worksheet('Data')
        (_write_bulk if bulk else _write_classic)(sheet, columns, rows)
        if extra is not None:
            extra(book, sheet)
        book.close()

    def assert_same(self, columns, rows, *, extra=None, prelude=None):
        with tempfile.TemporaryDirectory() as tmp:
            classic = Path(tmp) / 'classic.xlsx'
            bulk = Path(tmp) / 'bulk.xlsx'
            self.build(classic, bulk=False, columns=columns, rows=rows, extra=extra, prelude=prelude)
            self.build(bulk, bulk=True, columns=columns, rows=rows, extra=extra, prelude=prelude)
            classic_parts, bulk_parts = _read_parts(classic), _read_parts(bulk)
        self.assertEqual(sorted(classic_parts), sorted(bulk_parts))
        for name in classic_parts:
            self.assertEqual(classic_parts[name], bulk_parts[name], name)
        return classic_parts


class TestBulkWorksheetEquivalence(BulkEquivalenceCase):
    COLUMNS = ['id', 'name', 'value', 'note', 'empty', 'flag']

    def test_mixed_columns_blanks_and_literal_text(self):
        rows = [
            (1, 'alpha', 1.5, '=1+1', None, 'x'),
            (2, 'beta', float('nan'), None, None, '+SUM(A1)'),
            (3, 'alpha', -0.0, '@cmd', None, ''),
            (4, None, 2.25e-10, 'https://example.test/a', None, 'x'),
            (5, 'gamma', 12345678.901234567, '{=A1}', None, None),
            (2**60, 'delta', 1e300, '-2+3', None, 'long ' * 10),
        ]
        parts = self.assert_same(self.COLUMNS, rows)
        sheet_xml = parts['xl/worksheets/sheet1.xml'].decode('utf-8')
        self.assertNotIn('<f>', sheet_xml)
        self.assertNotIn('hyperlink', sheet_xml)

    def test_long_strings_are_truncated_like_write_string(self):
        rows = [(1, 'a' * 40000, 1.0, 'a' * 40000, None, 'b' * 32767)]
        self.assert_same(self.COLUMNS, rows)

    def test_shared_string_order_is_row_major_across_columns(self):
        rows = [
            (1, 'b', 1.0, 'a', None, 'c'),
            (2, 'a', 2.0, 'z', None, 'b'),
            (3, 'c', 3.0, 'b', None, 'q'),
        ]
        self.assert_same(self.COLUMNS, rows)

    def test_sparse_first_and_last_columns_change_row_spans(self):
        rows = []
        for index in range(40):
            first = 'f' if index % 7 == 0 else None
            middle = float(index)
            last = 'l' if index % 5 == 0 else None
            rows.append((first, middle, last))
        self.assert_same(['a', 'b', 'c'], rows)

    def test_fully_blank_rows_and_columns_write_nothing(self):
        rows = [(None, None, None), (None, 1.0, None), (None, None, None)] * 10
        self.assert_same(['a', 'b', 'c'], rows)

    def test_randomised_tables_match_cell_by_cell(self):
        generator = random.Random(1124)
        for table_index in range(6):
            width = generator.randint(1, 9)
            height = generator.randint(1, 70)
            kinds = [generator.choice(['text', 'float', 'int', 'mixednum']) for _ in range(width)]
            pool = ['', 'x', 'y z', '=A1', ' lead', 'trail ', 'ünï', 'a&b<c>', '"q"']
            rows = []
            for _ in range(height):
                row = []
                for kind in kinds:
                    if generator.random() < 0.2:
                        row.append(None)
                    elif kind == 'text':
                        row.append(generator.choice(pool))
                    elif kind == 'float':
                        row.append(generator.choice([0.0, -1.5, 3.14159, 1e-7, 123456789.125, float('nan')]))
                    elif kind == 'int':
                        row.append(generator.randint(-5, 10**12))
                    else:
                        row.append(generator.choice([1, 2.5, 3, 1e22]))
                rows.append(tuple(row))
            with self.subTest(table=table_index):
                self.assert_same([f'c{i}' for i in range(width)], rows)

    def test_column_format_set_after_the_bulk_write_is_applied(self):
        rows = [(1, 'a', 1.5), (2, 'b', 2.5), (3, 'c', None)]

        def extra(book, sheet):
            sheet.set_column(0, 0, 14, book.add_format({'align': 'center'}))
            sheet.set_column(1, 2, 20)
            sheet.autofilter(0, 0, 3, 2)
            sheet.freeze_panes(1, 0)

        self.assert_same(['a', 'b', 'c'], rows, extra=extra)

    def test_later_cell_writes_into_bulk_rows_win_and_keep_output_identical(self):
        rows = [(1, 'a'), (2, 'b'), (3, 'c'), (4, 'd')]

        def extra(book, sheet):
            sheet.write(2, 1, 'overwritten')
            sheet.write_number(3, 0, 99.5)
            sheet.write_url(1, 1, 'https://example.test/x', string='link')

        self.assert_same(['a', 'b'], rows, extra=extra)

    def test_row_level_formatting_inside_the_block_is_honoured(self):
        rows = [(1.0, 'a'), (2.0, 'b'), (3.0, 'c')]

        def extra(book, sheet):
            sheet.set_row(2, 30, book.add_format({'bold': True}))

        self.assert_same(['a', 'b'], rows, extra=extra)

    def test_chart_over_a_bulk_range_has_identical_cached_data(self):
        rows = [(float(index), f'cat{index % 3}') for index in range(12)]

        def extra(book, sheet):
            chart = book.add_chart({'type': 'line'})
            chart.add_series({'categories': '=Data!$B$2:$B$13', 'values': '=Data!$A$2:$A$13'})
            sheet.insert_chart('E2', chart)

        self.assert_same(['v', 'cat'], rows, extra=extra)

    def test_default_row_height_changes_force_the_classic_row_markup(self):
        rows = [(1.0, 'a'), (2.0, 'b')]

        def extra(book, sheet):
            sheet.set_default_row(24)

        self.assert_same(['a', 'b'], rows, extra=extra)

    def test_two_bulk_sheets_share_one_string_table(self):
        rows = [(1.0, 'shared'), (2.0, 'unique-a')]

        def extra(book, sheet):
            second = book.add_worksheet('Second')
            writer = _write_bulk if isinstance(second, BulkRowsWorksheet) else _write_classic
            writer(second, ['x', 'y'], [(5.0, 'shared'), (6.0, 'unique-b')])

        self.assert_same(['a', 'b'], rows, extra=extra)


class TestBulkSpans(unittest.TestCase):
    def test_large_block_span_work_scales_linearly(self):
        examined = []

        class CountedArray(np.ndarray):
            def __getitem__(self, key):
                result = super().__getitem__(key)
                # Boolean indexing must scan the entire mask, even for 16 hits.
                examined.append(key.size if isinstance(key, np.ndarray) else result.size)
                return result

        work = []
        for n_rows in (16_384, 65_536):
            sheet = BulkRowsWorksheet()
            block = xlsx_bulk_sheet.BulkBlock(first_row=7, n_rows=n_rows, columns=[])
            lows = np.zeros(n_rows, dtype=np.int64).view(CountedArray)
            highs = np.full(n_rows, 3, dtype=np.int64).view(CountedArray)
            examined.clear()
            with patch.object(sheet, '_row_column_extent', return_value=(lows, highs)):
                sheet._merge_block_spans(block)
            self.assertEqual(len(sheet.row_spans), block.last_row // 16 + 1)
            self.assertEqual(set(sheet.row_spans.values()), {'1:4'})
            work.append(sum(examined))
        self.assertLessEqual(work[0], 2 * 16_384)
        self.assertLessEqual(work[1], 4 * work[0])

    def test_sparse_unaligned_blocks_merge_with_existing_spans_like_stock(self):
        rows = [(None, None, None)] * 41
        rows[0] = (1.25, None, None)
        rows[40] = (None, None, 9.5)
        plan = plan_bulk_rows(['a', 'b', 'c'], rows)
        stock, bulk = xlsxwriter.worksheet.Worksheet(), BulkRowsWorksheet()
        stock.write_number(0, 5, 2)
        stock.write_number(50, 7, 1)
        for offset, row in enumerate(rows, start=7):
            for col, value in enumerate(row):
                write_untrusted_xlsx_cell(stock, offset, col, value)
        stock._calculate_spans()
        bulk.row_spans = {0: '6:6', 3: '8:8'}
        bulk._merge_block_spans(xlsx_bulk_sheet.BulkBlock(7, len(rows), plan.columns))
        self.assertEqual(bulk.row_spans, stock.row_spans)
        self.assertEqual(bulk.row_spans, {0: '1:6', 2: '3:3', 3: '8:8'})


class TestFastPathEquivalence(unittest.TestCase):
    """Cell, column and chart-cache fast paths versus stock XlsxWriter."""

    def assert_same_book(self, build):
        with tempfile.TemporaryDirectory() as tmp:
            parts = {}
            for label, fast in (('stock', False), ('fast', True)):
                path = Path(tmp) / f'{label}.xlsx'
                book = xlsxwriter.Workbook(str(path), xlsxwriter_workbook_options())
                if fast:
                    self.assertTrue(enable_bulk_rows(book))
                build(book)
                book.close()
                parts[label] = _read_parts(path)
        self.assertEqual(sorted(parts['stock']), sorted(parts['fast']))
        for name in parts['stock']:
            self.assertEqual(parts['stock'][name], parts['fast'][name], name)

    def test_sparse_unaligned_blocks_share_spans_with_cells_and_other_blocks(self):
        def build(book):
            sheet = book.add_worksheet('S')
            sheet.write_number(0, 5, 2)
            sheet.write_number(66, 7, 3)
            for first_row, height in ((7, 41), (49, 5), (56, 3), (64, 2)):
                rows = [(None, None, None)] * height
                rows[0] = ('=literal', None, None)
                rows[-1] = (None, None, 1.25)
                if isinstance(sheet, BulkRowsWorksheet):
                    self.assertTrue(sheet.write_bulk_block(first_row, plan_bulk_rows(['a', 'b', 'c'], rows)))
                else:
                    for offset, row in enumerate(rows, start=first_row):
                        for col, value in enumerate(row):
                            write_untrusted_xlsx_cell(sheet, offset, col, value)

        self.assert_same_book(build)

    def test_plain_columns_match_write_column(self):
        long_text = 'x' * 40000

        def build(book):
            sheet = book.add_worksheet('S')
            sheet.write_column(0, 0, [1.5, 2.25, -0.0, 1e-9, 12345678.123456789])
            sheet.write_column(0, 1, [1, 2, 3, 10**15])
            sheet.write_column(0, 2, ['a', 'b', 'a', 'ünï', 'a&b<c>', '_x0041_', ' lead '])
            sheet.write_column(0, 3, ['x', '', 'y'])  # empty text is a blank cell
            sheet.write_column(0, 4, ['{=A1}', 'ok', '{=SUM(A1)}'])  # `{=...}` is a formula for write()
            sheet.write_column(0, 5, [1.0, float('nan'), 3.0])
            sheet.write_column(0, 6, [1.0, float('inf'), 3.0])
            sheet.write_column(0, 7, [True, False])
            sheet.write_column(0, 8, [1.0, 'mixed', None])
            sheet.write_column(0, 9, [long_text, 'short'])  # truncated
            sheet.write_column('K1', [4.0, 5.0])  # A1 notation
            sheet.write_column(0, 11, (value for value in [7.0, 8.0]))  # generator input
            sheet.write_column(0, 12, [6.0, 7.0], book.add_format({'bold': True}))
            sheet.write_column(3, 13, [])
            sheet.write_column(2, 14, [9.0])

        self.assert_same_book(build)

    def test_numpy_scalars_use_the_stock_path(self):
        import numpy as np

        def build(book):
            sheet = book.add_worksheet('S')
            sheet.write_column(0, 0, [np.float64(1.5), np.float64(2.5)])
            sheet.write_column(0, 1, list(np.array([1, 2, 3])))
            sheet.write_column(0, 2, np.array([0.5, 0.25]))

        self.assert_same_book(build)

    def test_cell_markup_with_formats_row_formats_and_column_formats(self):
        def build(book):
            sheet = book.add_worksheet('S')
            bold = book.add_format({'bold': True})
            wrap = book.add_format({'text_wrap': True})
            sheet.write_number(0, 0, 1.5, bold)
            sheet.write_string(0, 1, 'a', wrap)
            sheet.write_number(1, 0, 2.5)
            sheet.write_string(1, 1, 'b')
            sheet.write_number(2, 0, 3)
            sheet.set_column(0, 0, 12, bold)
            sheet.set_column(1, 1, 20)
            sheet.set_row(1, 22, wrap)
            sheet.write(3, 0, True)
            sheet.write_blank(3, 1, None, bold)
            sheet.write_formula(4, 0, '=A1+A2', None, 4)
            sheet.write_datetime(4, 1, __import__('datetime').datetime(2024, 1, 2, 3, 4, 5), book.add_format({'num_format': 'yyyy-mm-dd'}))
            sheet.write_rich_string(5, 0, bold, 'rich ', 'text')
            sheet.write_url(5, 1, 'https://example.test/', string='link')
            sheet.write_column(6, 0, [1.0, 2.0, 3.0])
            sheet.write_column(6, 1, ['p', 'q', 'p'])

        self.assert_same_book(build)

    def test_chart_caches_match_stock_for_numeric_text_and_odd_tokens(self):
        def build(book):
            sheet = book.add_worksheet('S')
            sheet.write_column(0, 0, [1.0, 2.5, 1e-7, 30.0, 4.0])
            sheet.write_column(0, 1, ['2024-01-01', 'a&b', '<x>', '_x0041_', 'ünï'])
            sheet.write_column(0, 2, [1, 2, 3, 4, 5])
            sheet.write(2, 3, 'text-in-numeric-range')
            sheet.write(1, 3, 9.5)
            sheet.write_formula(3, 3, '=A1', None, 7)
            for categories in ('=S!$B$1:$B$5', '=S!$C$1:$C$5', None):
                chart = book.add_chart({'type': 'line'})
                series = {'values': '=S!$A$1:$A$5'}
                if categories:
                    series['categories'] = categories
                chart.add_series(series)
                chart.add_series({'values': '=S!$D$1:$D$5', 'categories': '=S!$B$1:$B$5'})
                sheet.insert_chart('F2', chart)
            scatter = book.add_chart({'type': 'scatter'})
            scatter.add_series({'categories': '=S!$A$1:$A$5', 'values': '=S!$C$1:$C$5'})
            sheet.insert_chart('F20', scatter)

        self.assert_same_book(build)


class TestBulkPlanning(unittest.TestCase):
    def test_unsupported_values_use_the_per_cell_path(self):
        class Marker(float):
            pass

        columns = ['a', 'b']
        self.assertIsNone(plan_bulk_rows(columns, [(True, 'x')]))
        self.assertIsNone(plan_bulk_rows(columns, [(float('inf'), 'x')]))
        self.assertIsNone(plan_bulk_rows(columns, [(float('-inf'), 'x')]))
        self.assertIsNone(plan_bulk_rows(columns, [(1.0, 'x'), ('text', 'y')]))
        self.assertIsNone(plan_bulk_rows(columns, [(Marker(1.0), 'x')]))
        self.assertIsNone(plan_bulk_rows(columns, [(10**400, 'x')]))
        self.assertIsNone(plan_bulk_rows(columns, [(1.0, b'bytes')]))
        self.assertIsNone(plan_bulk_rows(columns, [(1.0, 'x'), (2.0,)]))
        self.assertIsNone(plan_bulk_rows(columns, []))
        self.assertIsNone(plan_bulk_rows([], [()]))
        self.assertIsNotNone(plan_bulk_rows(columns, [(1, 'x'), (None, None), (float('nan'), 'y')]))

    def test_unsupported_tables_still_produce_literal_output_via_the_backend(self):
        from metroliza.exporting.export_backends import ExcelExportBackend

        backend = ExcelExportBackend()
        with tempfile.TemporaryDirectory() as tmp:
            outputs = {}
            for label, flag in (('bulk', '1'), ('classic', '0')):
                with patch.dict('os.environ', {BULK_WRITE_ENV: flag}):
                    writer = backend.create_writer(str(Path(tmp) / f'{label}.xlsx'))
                    backend.write_dataframe(
                        writer,
                        ([('=1+1', float('inf'), True, '@cmd'), ('+2', 1.0, False, 'https://x.test')], ['a', 'b', 'c', 'd']),
                        'Data',
                    )
                    backend.close_writer(writer)
                outputs[label] = _read_parts(Path(tmp) / f'{label}.xlsx')
        self.assertEqual(outputs['bulk'], outputs['classic'])


class TestBulkSwitches(unittest.TestCase):
    def test_environment_switch_disables_the_bulk_path(self):
        with patch.dict('os.environ', {BULK_WRITE_ENV: '0'}):
            self.assertFalse(bulk_write_enabled())
            book = types.SimpleNamespace(worksheet_class=xlsxwriter.worksheet.Worksheet)
            self.assertFalse(enable_bulk_rows(book))
            self.assertIsNot(book.worksheet_class, BulkRowsWorksheet)
        with patch.dict('os.environ', {BULK_WRITE_ENV: '1'}):
            self.assertTrue(bulk_write_enabled())

    def test_unverified_xlsxwriter_versions_use_the_classic_path(self):
        with patch.object(xlsx_bulk_sheet.xlsxwriter, '__version__', '3.9.0'):
            self.assertFalse(bulk_write_enabled())
        with patch.object(xlsx_bulk_sheet.xlsxwriter, '__version__', 'weird'):
            self.assertFalse(bulk_write_enabled())

    def test_numeric_values_round_trip_like_the_classic_writer(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'numbers.xlsx'
            book = xlsxwriter.Workbook(str(path), xlsxwriter_workbook_options())
            enable_bulk_rows(book)
            sheet = book.add_worksheet('Data')
            values = [0.1 + 0.2, 1 / 3, 2**53 + 1, -1e-300]
            _write_bulk(sheet, ['v'], [(value,) for value in values])
            book.close()
            import openpyxl

            read_back = [cell.value for cell in openpyxl.load_workbook(path)['Data']['A'][1:]]
        for original, stored in zip(values, read_back):
            self.assertTrue(math.isclose(float(original), stored, rel_tol=1e-15))


if __name__ == '__main__':
    unittest.main()
