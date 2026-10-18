"""Workbook-equivalence guard for the default-export write path (#1124).

The default export writes the raw ``MEASUREMENTS`` sheet and the per-reference
measurement sheets through bulk/vectorised paths.  These tests export one fixed
synthetic database twice -- once through the optimised path and once with the
optimised path disabled (``METROLIZA_EXPORT_BULK_WRITE=0``, the original
cell-by-cell ``write_untrusted_xlsx_cell`` semantics) -- and require identical
workbooks: package parts, openpyxl read-back (sheet names, cell values/types,
number formats), and charts (count, type, series ranges and titles).

The literal-text policy of ``write_untrusted_xlsx_cell`` (formula-like and
URL-like text stays text) is asserted explicitly on the optimised output.
"""

import hashlib
import importlib.util
import math
import re
import sys
import tempfile
import types
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import xlsxwriter

from metroliza.charts.matplotlib_runtime import configure_headless_matplotlib
from modules.contracts import AppPaths, ExportOptions, ExportRequest  # noqa: E402
from modules.report_repository import ReportRepository  # noqa: E402
from modules.report_schema import ensure_report_schema  # noqa: E402

BULK_WRITE_ENV = 'METROLIZA_EXPORT_BULK_WRITE'

# Strings that spreadsheet applications would otherwise interpret.
HOSTILE_TEXT = {
    'part_name': '=HYPERLINK("http://example.test/p","click")',
    'revision': '+1+1',
    'operator_name': '@SUM(A1:A2)',
    'description': '{=A1}',
    'feature_label': '-2+3',
    'section_name': 'https://example.test/section?a=1',
}
LONG_TEXT = 'L' * 33000  # longer than the 32767-character cell limit
HOSTILE_FILE_NAME = '=cmd_report.pdf'

_VOLATILE_PARTS = {'docProps/core.xml'}
_NS_MAIN = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


class _DummyThread:
    def __init__(self, *args, **kwargs):
        pass


class _DummyCoreApp:
    @staticmethod
    def processEvents():
        return None


def _dummy_signal(*args, **kwargs):
    class _Signal:
        def emit(self, *a, **k):
            return None

    return _Signal()


class _DummyLogger:
    def __init__(self, *args, **kwargs):
        pass


def _load_export_thread_type():
    """Load ``ExportDataThread`` without a Qt event loop (see test_export_workbook_output)."""
    configure_headless_matplotlib()
    for dependency_name in ('numpy', 'pandas', 'matplotlib.pyplot', 'scipy.stats', 'seaborn'):
        importlib.import_module(dependency_name)
    qtcore_stub = types.ModuleType('PyQt6.QtCore')
    qtcore_stub.QCoreApplication = _DummyCoreApp
    qtcore_stub.QThread = _DummyThread
    qtcore_stub.pyqtSignal = _dummy_signal

    custom_logger_stub = types.ModuleType('modules.custom_logger')
    custom_logger_stub.CustomLogger = _DummyLogger

    module_name = '_test_export_data_thread_workbook_equivalence'
    module_path = (
        Path(__file__).resolve().parents[1] / 'src' / 'metroliza' / 'exporting' / 'export_data_thread.py'
    )
    with patch.dict(
        sys.modules,
        {'PyQt6.QtCore': qtcore_stub, 'modules.custom_logger': custom_logger_stub},
    ):
        spec = importlib.util.spec_from_file_location(module_name, module_path)
        export_module = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = export_module
        try:
            spec.loader.exec_module(export_module)
        finally:
            sys.modules.pop(module_name, None)
    return export_module.ExportDataThread


def _measurement(header, ax, nominal, tol_plus, tol_minus, meas, dev, row_order, **text):
    return {
        'page_number': 1,
        'row_order': row_order,
        'header': header,
        'section_name': text.get('section_name', header),
        'feature_label': text.get('feature_label', header),
        'characteristic_name': 'LOC',
        'characteristic_family': 'LOC',
        'description': text.get('description', f'{header} LOC'),
        'ax': ax,
        'nominal': nominal,
        'tol_plus': tol_plus,
        'tol_minus': tol_minus,
        'bonus': 0.0,
        'meas': meas,
        'dev': dev,
        'outtol': 0,
        'is_nok': False,
        'status_code': 'ok',
    }


def build_equivalence_db(db_path, work_dir, *, references=('REF-A', 'REF-B', 'REF=SUM(1)'), reports_per_reference=5):
    """Create a small deterministic synthetic report database (no real data)."""
    work_dir = Path(work_dir)
    source_dir = work_dir / 'reports'
    source_dir.mkdir(parents=True, exist_ok=True)
    ensure_report_schema(str(db_path))
    repository = ReportRepository(str(db_path))

    report_id = 0
    for ref_index, reference in enumerate(references):
        for sample in range(1, reports_per_reference + 1):
            report_id += 1
            file_name = f'part_{report_id}.pdf'
            if report_id == 2:
                file_name = HOSTILE_FILE_NAME
            source_path = source_dir / file_name
            source_path.write_bytes(f'synthetic-report-{report_id}'.encode('utf-8'))
            base = 10.0 * (ref_index + 1)
            measurements = []
            for feature in range(1, 4):
                header = f'FEATURE_{feature}'
                for ax_index, ax in enumerate(('X', 'Y')):
                    nominal = base + feature + ax_index
                    drift = ((sample * 7 + feature * 3 + ax_index) % 11 - 5) / 40.0
                    text = {}
                    if report_id == 3 and feature == 1 and ax == 'X':
                        text = dict(HOSTILE_TEXT)
                    if report_id == 4 and feature == 2 and ax == 'Y':
                        text = {'description': LONG_TEXT}
                    dev = round(drift, 4)
                    if report_id == 5 and feature == 3 and ax == 'X':
                        dev = None  # NULL numeric value -> blank cell
                    if feature == 3 and ax == 'Y':
                        # One-sided geometric tolerance (nominal 0, lower limit 0).
                        measurements.append(
                            _measurement(header, ax, 0.0, 0.25, 0.0, round(abs(drift), 4), dev, len(measurements) + 1, **text)
                        )
                    else:
                        measurements.append(
                            _measurement(
                                header, ax, nominal, 0.5, -0.5, round(nominal + drift, 4), dev,
                                len(measurements) + 1, **text,
                            )
                        )
            metadata = {
                'reference': reference,
                'reference_raw': reference,
                'report_date': f'2024-02-{sample:02d}',
                'part_name': HOSTILE_TEXT['part_name'] if report_id == 3 else 'Part',
                'revision': HOSTILE_TEXT['revision'] if report_id == 3 else 'A',
                'sample_number': str(sample),
                'sample_number_kind': 'explicit_sample_number',
                'stats_count_raw': str(sample),
                'stats_count_int': sample,
                'operator_name': HOSTILE_TEXT['operator_name'] if report_id == 3 else 'Operator',
                'comment': None,
            }
            repository.persist_parsed_report(
                source_path=source_path,
                parser_id='cmm',
                parser_version='1.0',
                template_family='cmm_pdf_header_box',
                template_variant='cmm_pdf_header_box_serial_variant',
                parse_status='parsed_with_warnings',
                metadata=metadata,
                measurements=measurements,
                candidates=[],
                warnings=[],
                metadata_version='report_metadata_v1',
                metadata_profile_id='cmm_pdf_header_box',
                metadata_profile_version='1',
                page_count=1,
                measurement_count=len(measurements),
                has_nok=False,
                nok_count=0,
                metadata_confidence=0.9,
                identity_hash=f'identity-{report_id}',
            )
    return report_id


def export_workbook(db_path, out_path, *, bulk, export_thread_type, options=None):
    """Run the real export backend against ``db_path`` and return ``out_path``."""
    options = options if options is not None else ExportOptions(generate_summary_sheet=False)
    request = ExportRequest(paths=AppPaths(db_file=str(db_path), excel_file=str(out_path)), options=options)
    env_value = '1' if bulk else '0'
    with patch.dict('os.environ', {BULK_WRITE_ENV: env_value}):
        thread = export_thread_type(request)
        if not thread.get_export_backend().run(thread):
            raise AssertionError('export did not complete')
    return Path(out_path)


def read_package_parts(path):
    """Return ``{part_name: bytes}`` with volatile parts normalised."""
    parts = {}
    with zipfile.ZipFile(path, 'r') as workbook_zip:
        for name in workbook_zip.namelist():
            data = workbook_zip.read(name)
            if name in _VOLATILE_PARTS:
                data = re.sub(rb'<dcterms:(created|modified)[^>]*>[^<]*</dcterms:\1>', b'', data)
            parts[name] = data
    return parts


def describe_workbook(path):
    """Normalised openpyxl read-back: sheets, cells (value/type/format) and charts."""
    import openpyxl

    workbook = openpyxl.load_workbook(path)
    description = {}
    for worksheet in workbook.worksheets:
        cells = {}
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None and cell.data_type == 'n' and cell.number_format == 'General':
                    continue
                cells[cell.coordinate] = (cell.value, cell.data_type, cell.number_format)
        charts = []
        for chart in worksheet._charts:
            series = []
            for serie in chart.series:
                title = None
                if serie.tx is not None and serie.tx.strRef is not None:
                    title = serie.tx.strRef.f
                elif serie.tx is not None:
                    title = serie.tx.v
                series.append(
                    (
                        title,
                        serie.cat.numRef.f if serie.cat is not None and serie.cat.numRef is not None
                        else (serie.cat.strRef.f if serie.cat is not None and serie.cat.strRef is not None else None),
                        serie.val.numRef.f if serie.val is not None and serie.val.numRef is not None else None,
                        serie.xVal.numRef.f if getattr(serie, 'xVal', None) is not None and serie.xVal.numRef is not None else None,
                        serie.yVal.numRef.f if getattr(serie, 'yVal', None) is not None and serie.yVal.numRef is not None else None,
                    )
                )
            charts.append((chart.tagname, tuple(series)))
        description[worksheet.title] = {
            'dimensions': worksheet.dimensions,
            'freeze_panes': worksheet.freeze_panes,
            'auto_filter': worksheet.auto_filter.ref,
            'cells': cells,
            'charts': charts,
            'hyperlinks': sorted(str(link.ref) for link in worksheet._hyperlinks),
        }
    return description


def sheet_part_paths(path):
    """Map worksheet name -> package part path."""
    import xml.etree.ElementTree as ET

    with zipfile.ZipFile(path, 'r') as workbook_zip:
        workbook_xml = ET.fromstring(workbook_zip.read('xl/workbook.xml'))
        rels = ET.fromstring(workbook_zip.read('xl/_rels/workbook.xml.rels'))
    rel_targets = {rel.attrib['Id']: rel.attrib['Target'] for rel in rels}
    rid_key = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    return {
        sheet.attrib['name']: 'xl/' + rel_targets[sheet.attrib[rid_key]].lstrip('/').removeprefix('xl/')
        for sheet in workbook_xml.iter(f'{_NS_MAIN}sheet')
    }


def assert_workbooks_equivalent(testcase, reference_path, candidate_path):
    reference_parts = read_package_parts(reference_path)
    candidate_parts = read_package_parts(candidate_path)
    testcase.assertEqual(sorted(reference_parts), sorted(candidate_parts))
    for name in sorted(reference_parts):
        if reference_parts[name] != candidate_parts[name]:
            testcase.fail(
                f'package part {name} differs: '
                f'{hashlib.sha256(reference_parts[name]).hexdigest()[:12]} != '
                f'{hashlib.sha256(candidate_parts[name]).hexdigest()[:12]}'
            )
    reference_view = describe_workbook(reference_path)
    candidate_view = describe_workbook(candidate_path)
    testcase.assertEqual(list(reference_view), list(candidate_view))
    for sheet_name, reference_sheet in reference_view.items():
        testcase.assertEqual(reference_sheet, candidate_view[sheet_name], sheet_name)


class TestWorkbookEquivalenceHelpers(unittest.TestCase):
    def test_xlsxwriter_options_keep_text_literal(self):
        from metroliza.exporting.xlsx_writer_policy import xlsxwriter_workbook_options

        options = xlsxwriter_workbook_options()
        self.assertFalse(options['strings_to_formulas'])
        self.assertFalse(options['strings_to_urls'])
        self.assertTrue(options['nan_inf_to_errors'])
        self.assertTrue(hasattr(xlsxwriter, '__version__'))


class TestDefaultExportWorkbookEquivalence(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.export_thread_type = _load_export_thread_type()
        cls._tmp = tempfile.TemporaryDirectory()
        tmp = Path(cls._tmp.name)
        cls.db_path = tmp / 'metroliza.sqlite'
        cls.total_reports = build_equivalence_db(cls.db_path, tmp)
        cls.bulk_calls = {'reference': [], 'candidate': []}
        cls.reference_path = cls._export(tmp / 'reference.xlsx', bulk=False, label='reference')
        cls.candidate_path = cls._export(tmp / 'candidate.xlsx', bulk=True, label='candidate')

    @classmethod
    def _export(cls, out_path, *, bulk, label):
        from metroliza.exporting.xlsx_bulk_sheet import BulkRowsWorksheet

        original = BulkRowsWorksheet.write_bulk_block

        def recording(worksheet, first_row, plan):
            accepted = original(worksheet, first_row, plan)
            cls.bulk_calls[label].append((worksheet.name, plan.n_rows, accepted))
            return accepted

        with patch.object(BulkRowsWorksheet, 'write_bulk_block', recording):
            return export_workbook(cls.db_path, out_path, bulk=bulk, export_thread_type=cls.export_thread_type)

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_optimised_run_used_the_bulk_path_and_reference_run_did_not(self):
        self.assertEqual(self.bulk_calls['reference'], [])
        self.assertEqual(self.bulk_calls['candidate'], [('MEASUREMENTS', self.total_reports * 6, True)])

    def test_default_export_workbooks_are_identical(self):
        assert_workbooks_equivalent(self, self.reference_path, self.candidate_path)

    def test_workbook_structure_is_the_expected_default_export(self):
        view = describe_workbook(self.candidate_path)
        self.assertEqual(list(view)[-1], 'MEASUREMENTS')
        self.assertGreaterEqual(len(view), 4)
        total_charts = sum(len(sheet['charts']) for sheet in view.values())
        self.assertGreaterEqual(total_charts, 9)
        raw = view['MEASUREMENTS']
        # 3 references x 5 reports x 3 features x 2 axes, plus the header row.
        self.assertEqual(raw['dimensions'].split(':')[0], 'A1')
        self.assertTrue(raw['dimensions'].endswith(str(1 + self.total_reports * 6)))
        self.assertEqual(raw['freeze_panes'], 'A2')
        self.assertTrue(raw['auto_filter'].startswith('A1:'))

    def test_untrusted_text_stays_literal(self):
        import openpyxl

        workbook = openpyxl.load_workbook(self.candidate_path)
        seen = {text: 0 for text in (*HOSTILE_TEXT.values(), HOSTILE_FILE_NAME, LONG_TEXT[:32767])}
        for worksheet in workbook.worksheets:
            self.assertEqual(list(worksheet._hyperlinks), [], worksheet.title)
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == 'f' and worksheet.title == 'MEASUREMENTS':
                        self.fail(f'unexpected formula-typed cell {worksheet.title}!{cell.coordinate}')
                    if isinstance(cell.value, str) and cell.value in seen:
                        seen[cell.value] += 1
                        self.assertEqual(cell.data_type, 's', f'{worksheet.title}!{cell.coordinate}')
        raw_sheet = workbook['MEASUREMENTS']
        raw_values = {cell.value for row in raw_sheet.iter_rows() for cell in row if isinstance(cell.value, str)}
        for text in (*HOSTILE_TEXT.values(), HOSTILE_FILE_NAME, LONG_TEXT[:32767]):
            self.assertIn(text, raw_values)

    def test_raw_sheet_numbers_blanks_and_types(self):
        import openpyxl

        reference = openpyxl.load_workbook(self.reference_path)['MEASUREMENTS']
        candidate = openpyxl.load_workbook(self.candidate_path)['MEASUREMENTS']
        headers = [cell.value for cell in candidate[1]]
        self.assertIn('meas', [str(h).lower() for h in headers])
        numeric_cells = 0
        blank_cells = 0
        for ref_row, cand_row in zip(reference.iter_rows(), candidate.iter_rows()):
            for ref_cell, cand_cell in zip(ref_row, cand_row):
                self.assertEqual((ref_cell.value, ref_cell.data_type), (cand_cell.value, cand_cell.data_type))
                if isinstance(cand_cell.value, float):
                    numeric_cells += 1
                    self.assertTrue(math.isfinite(cand_cell.value))
                if cand_cell.value is None:
                    blank_cells += 1
        self.assertGreater(numeric_cells, 100)
        self.assertGreater(blank_cells, 0)


if __name__ == '__main__':
    unittest.main()
