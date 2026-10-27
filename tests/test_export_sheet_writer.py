import unittest
from unittest.mock import patch
import weakref
from defusedxml import ElementTree
from zipfile import ZipFile

import pandas as pd

from modules.export_sheet_writer import (
    build_measurement_header_block_plan,
    build_measurement_summary_row_layout,
    build_measurement_write_bundle,
    build_measurement_write_bundle_cached,
    build_summary_panel_write_plan,
    write_measurement_block,
    write_measurement_summary_rows,
)
import xlsxwriter
import tempfile


# Ordered package fingerprints of valid measurement workbooks from the pre-fix source.
VALID_MEASUREMENT_XLSX_FINGERPRINTS = (
    'aa3b49fa5d5c7a793ca13149e170428b6a05ffbd588f3d01f9794ab64943f302',
    'dbcaf74a9174a04813b47d460e0e2ee912e422fbe41c1cc591af270e45b204c3',
    'c0d087a154a1703c806f7ca77483ad45b2304851cc7a4ebe7399497a2ff76c3c',
    '85e82cead80d6098c3c76c7a71b0740ee1414cd2dfbcb4ffe49352547b99b6ab',
    'c7331f964e89f5810703749af0762ac646e59b6ba6020c84d8e5555736713c2c',
    '1ab35267e9da678a1dc4cf16da5c4f8e0b954013218a2ceae18cc6e79997bc52',
    'b397933785cebb212a6593e05bcac2769a0af328d22605393b73951a1fdbdf7c',
    '279bee43a07dbcceaf745e1c4640b020971087822a80b72ffc99bf9abf0228af',
    'ba167496b7511f3233df6ba16becdab20774627b85d116a16f97aeb4c4bedebd',
    '2b52aef4c684c4e0da18940a3e2c37d8d940f74a74746ffda2f3e90a84e3e177',
    '130ead47781396b4aa571aa065d206eeaff0cf0fab23eaf0d353fe71ab3cdecf',
    'd9cc19dcc1f01456c43d095330d848366ef58a14815db0c2c1b9b80cd581a264',
    '16c35f82483f7a00e92368c3e909033f4c9d3ddddd18f3ca58cc303304f09419',
    'f323641b811c951a5419cb42abc23086ff62b436f1e6be1cbb3eb5f2eb84efcc',
    '79603c746e554577039926010421ac278e6719e7b680eab676af54a34ebd10ea',
    '6037cf023897618f56eeb43b8de8ab586f9f60ec856a385d6ac64e5f12620801',
    '757c957096414dcdebdad28e58dcf7bf7ec93d42a5e6d8ce9b795c18afc802db',
    '1ba12b2339985bd6ad37cc52af47267e37a6e02e0286658a4c18a32fac201c45',
    '3855e920e59c2a24c58dfc7177c71c146b4e2127b5ff284436972c0af2e436e6',
    'ca168680698f60aa72026aa96569c9e8e6b4c64b934802bd5489cdd947c183e7',
    '929a0ce7cea8a97ba585a38bf9e9909076df3863516d42718dd50c2da6558e53',
    '128d71c3fdf883e013d1eff770b0d30b97726a69752f7990454f32ffd69a4031',
    '082c6c55fe9b303380ae3bfa511a7e486bd940f2cde011bdd32b634f921be297',
    'a033b7849eea4e5172c75d13eceede55ef1910590e0a1279baf237571fbe1107',
    'dda1e914a804be750462c6baf1133b2a883b3ceb186c1b52141259800a515fc4',
    '45379a1f9ffd0122de05b04afa95de7d8e7a5d942200736db0505918de67e023',
    '1ff84c01f1319124e0f9f410fcf931b25543a758039e9a5cd13b4ed98f5dfe97',
    'fa72da13afdb75edfd3ab0f93027465d4eb165166d9de5ab06f928a64fdf5b81',
    '26082de1b7a44a365ebc3627be59d6e81c703478c35b22ffe25c8ace1f1f2b80',
    '4bffa2fffa0c026ffde4f6ed9b7c1a03d96154f20d2d4bfc95aff3255ea994d4',
    '5d0053cb742179d7367d766b8931cd50dabb756611fd687bf680dd25986d1b76',
    '213c8a2449b29f293bb3db2c89fbbfe0f0169a0ebfe2468c575fad2dc3818bdb',
    'd2c2254aa9cf2a35838fe5d2e28a90b605800a98ed85f3a71413517745cf2ce7',
    'd2c68d811e4090cbf0b5fb5e659a7b5a2ceca998feeb2feff146202088fab6fb',
    '6814db8d360cfd310c462df093bdaedeec1419e7466acefd853c9ecba167b508',
)


def _measurement_workbook_fingerprint(path):
    import hashlib
    import json
    from zipfile import ZipFile
    import zlib
    from scripts.compare_export_outputs import _core_without_times

    digest = hashlib.sha256()
    with ZipFile(path) as package:
        for member in package.infolist():
            data = package.read(member.filename)
            crc = member.CRC
            if member.filename == 'docProps/core.xml':
                data = _core_without_times(data)
                crc = zlib.crc32(data)
            digest.update(json.dumps((member.filename, crc, len(data))).encode() + b'\0' + data)
    return digest.hexdigest()


class DummyWorksheet:
    def __init__(self):
        self.writes = []
        self.formulas = []
        self.columns = []
        self.conditional_formats = []

    def write(self, row, col, value, *args, **kwargs):
        self.writes.append((row, col, value))

    def write_formula(self, row, col, formula, *args, **kwargs):
        self.formulas.append((row, col, formula))

    def write_column(self, row, col, values):
        self.columns.append((row, col, list(values)))

    def conditional_format(self, *args):
        self.conditional_formats.append(args)


class TestExportSheetWriter(unittest.TestCase):
    def _write_valid_measurement_corpus_case(self, path, index, owners):
        from metroliza.exporting.export_backends import ExcelExportBackend
        from metroliza.exporting.export_sheet_writer import create_measurement_formats
        from openpyxl import load_workbook

        nominal = (0.0, 10.0, -3.0, 1000000.0, 0.125)[index % 5]
        tolerance = (0.5, 0.125, 2.0, 5.0, 0.005)[index % 5]
        row_count, block_count = 1 + index % 7, 1 + index % 3
        measurements = [round(nominal + (row % 5 - 2) * tolerance, 3) for row in range(row_count)]
        frame = pd.DataFrame({'DATE': [f'2026-01-{row + 1:02d}' for row in range(row_count)],
                              'SAMPLE_NUMBER': [f'S-{index}-{row}' for row in range(row_count)],
                              'MEAS': measurements, 'NOM': [nominal] * row_count,
                              '+TOL': [tolerance] * row_count, '-TOL': [-tolerance] * row_count})
        backend = ExcelExportBackend()
        writer = backend.create_writer(str(path))
        owner = backend.get_workbook(writer)
        # Keep distinct owners alive for this valid-output corpus; do not alter the cache or id().
        owners.append(owner)
        try:
            formats = create_measurement_formats(owner)
            worksheet = owner.add_worksheet('MEASUREMENTS')
            worksheet.set_column(0, block_count * 5 - 1, 12, formats['default'])
            for block in range(block_count):
                column = block * 5
                header = f'Feature {index} / Δ-{block}'
                bundle = build_measurement_write_bundle(header, frame, column)
                write_measurement_block(worksheet, bundle, formats, base_col=column)
                worksheet.set_column(column + 4, column + 4, None, formats['border'])
                chart = owner.add_chart({'type': 'line'})
                chart.add_series({'name': header,
                                  'categories': ['MEASUREMENTS', 21, column, 20 + row_count, column],
                                  'values': ['MEASUREMENTS', 21, column + 2, 20 + row_count, column + 2]})
                worksheet.insert_chart(7, column, chart)
        finally:
            backend.close_writer(writer)
        workbook = load_workbook(path)
        try:
            sheet = workbook['MEASUREMENTS']
            self.assertEqual(len(sheet._charts), block_count)
            for block in range(block_count):
                column = block * 5
                self.assertEqual(sheet.cell(1, column + 2).value, nominal)
                self.assertEqual(sheet.cell(21, column + 3).value, f'Feature {index} / Δ-{block}')
                self.assertEqual(sheet.cell(7, column + 2).number_format, '0.00%')
                self.assertTrue(sheet.cell(7, column + 2).value.startswith('='))
                self.assertEqual([sheet.cell(22 + row, column + 3).value for row in range(row_count)],
                                 measurements)
        finally:
            workbook.close()

    def test_valid_measurement_output_corpus_matches_pre_fix_packages(self):
        from pathlib import Path

        self.assertEqual(len(VALID_MEASUREMENT_XLSX_FINGERPRINTS), 35)
        owners = []
        with tempfile.TemporaryDirectory() as directory:
            for index, expected in enumerate(VALID_MEASUREMENT_XLSX_FINGERPRINTS):
                with self.subTest(case=index):
                    path = Path(directory) / f'case-{index:02d}.xlsx'
                    self._write_valid_measurement_corpus_case(path, index, owners)
                    self.assertEqual(_measurement_workbook_fingerprint(path), expected)

    def test_build_measurement_write_bundle_includes_limit_columns_contract(self):
        header_group = pd.DataFrame(
            {
                'DATE': ['2024-01-01', '2024-01-02', '2024-01-03'],
                'SAMPLE_NUMBER': ['1', '2', '3'],
                'MEAS': [10.1, 10.2, 10.3],
                'NOM': [10.0, 10.0, 10.0],
                '+TOL': [0.5, 0.5, 0.5],
                '-TOL': [-0.5, -0.5, -0.5],
            }
        )

        bundle = build_measurement_write_bundle('Diameter - X', header_group, 0)

        labels = [column[2] for column in bundle['data_columns']]
        self.assertEqual(labels, ['Date', 'Sample #', 'Diameter - X', 'USL', 'LSL'])

        usl_values = bundle['data_columns'][3][3]
        lsl_values = bundle['data_columns'][4][3]
        self.assertEqual(usl_values, [10.5, 10.5, 10.5])
        self.assertEqual(lsl_values, [9.5, 9.5, 9.5])
        self.assertTrue(all(value is not None for value in usl_values))
        self.assertTrue(all(value is not None for value in lsl_values))

        measurement_plan = bundle['measurement_plan']
        self.assertIn('usl_column', measurement_plan)
        self.assertIn('lsl_column', measurement_plan)

    def test_write_measurement_block_applies_three_conditional_rules(self):
        header_group = pd.DataFrame(
            {
                'DATE': ['2024-01-01', '2024-01-02'],
                'SAMPLE_NUMBER': ['1', '2'],
                'MEAS': [10.1, 10.2],
                'NOM': [10.0, 10.0],
                '+TOL': [0.5, 0.5],
                '-TOL': [-0.5, -0.5],
            }
        )
        bundle = build_measurement_write_bundle('Diameter - X', header_group, 0)
        worksheet = DummyWorksheet()
        formats = {'percent': object(), 'wrap': object(), 'red': object()}

        measurement_plan = write_measurement_block(worksheet, bundle, formats, base_col=0)

        self.assertEqual(measurement_plan['data_start_row'], 21)
        self.assertEqual(measurement_plan['chart_insert_row'], 7)
        self.assertEqual(len(worksheet.conditional_formats), 3)
        self.assertEqual(worksheet.conditional_formats[0][4]['value'], '=($B$1+$B$2)')
        self.assertEqual(worksheet.conditional_formats[1][4]['value'], '=($B$1+$B$3)')
        self.assertTrue(any(w[2] == 'NOK %' for w in worksheet.writes if isinstance(w[2], str)))
        self.assertTrue(any((w[0], w[1], w[2]) == (0, 2, 'MIN') for w in worksheet.writes))

    def test_write_measurement_block_skips_lower_conditional_rule_for_one_sided_gdt(self):
        header_group = pd.DataFrame(
            {
                'DATE': ['2024-01-01', '2024-01-02'],
                'SAMPLE_NUMBER': ['1', '2'],
                'MEAS': [0.02, 0.12],
                'NOM': [0.0, 0.0],
                '+TOL': [0.1, 0.1],
                '-TOL': [0.0, 0.0],
            }
        )
        bundle = build_measurement_write_bundle('A-V3 - TP', header_group, 0)
        worksheet = DummyWorksheet()
        formats = {'percent': object(), 'wrap': object(), 'red': object()}

        write_measurement_block(worksheet, bundle, formats, base_col=0)

        self.assertEqual(len(worksheet.conditional_formats), 2)
        self.assertEqual(worksheet.conditional_formats[0][4]['value'], '=($B$1+$B$2)')
        self.assertFalse(
            any(item[4].get('value') == '=($B$1+$B$3)' for item in worksheet.conditional_formats)
        )

    def test_build_measurement_summary_row_layout_keeps_legacy_coordinates(self):
        stat_rows = [
            ('MIN', '=MIN(C22:C30)', None),
            ('NOK %', '=10%', 'percent'),
        ]

        layout = build_measurement_summary_row_layout(base_col=6, stat_rows=stat_rows)

        self.assertEqual(layout[0]['row'], 3)
        self.assertEqual(layout[0]['label_col'], 6)
        self.assertEqual(layout[0]['value_col'], 7)
        self.assertEqual(layout[1]['row'], 4)
        self.assertEqual(layout[1]['style'], 'percent')

    def test_write_measurement_summary_rows_writes_formula_and_percent_style(self):
        worksheet = DummyWorksheet()
        summary_rows = [
            {'row': 3, 'label_col': 0, 'value_col': 1, 'label': 'MIN', 'formula': '=MIN(C22:C30)', 'style': None},
            {'row': 4, 'label_col': 0, 'value_col': 1, 'label': 'NOK %', 'formula': '=10%', 'style': 'percent'},
        ]

        write_measurement_summary_rows(worksheet, summary_rows, formats={'percent': object()})

        self.assertIn((3, 0, 'MIN'), worksheet.writes)
        self.assertIn((4, 0, 'NOK %'), worksheet.writes)
        self.assertEqual(worksheet.formulas[0], (3, 1, '=MIN(C22:C30)'))
        self.assertEqual(worksheet.formulas[1], (4, 1, '=10%'))


    def test_build_summary_panel_write_plan_maps_header_and_image_slots(self):
        anchors = {
            'header': (40, 0),
            'distribution': (41, 1),
            'iqr': (41, 11),
            'histogram': (41, 21),
            'trend': (41, 35),
        }

        plan = build_summary_panel_write_plan(anchors, 'DIA - X')

        self.assertEqual(plan['header_cell'], {'row': 40, 'col': 0, 'value': 'DIA - X'})
        self.assertEqual(plan['image_slots']['distribution'], {'row': 41, 'col': 1})
        self.assertEqual(plan['image_slots']['iqr'], {'row': 41, 'col': 11})
        self.assertEqual(plan['image_slots']['histogram'], {'row': 41, 'col': 21})
        self.assertEqual(plan['image_slots']['trend'], {'row': 41, 'col': 35})


    def test_cached_write_bundle_matches_uncached_contract(self):
        header_group = pd.DataFrame(
            {
                'DATE': ['2024-01-01', '2024-01-02'],
                'SAMPLE_NUMBER': ['1', '2'],
                'MEAS': [10.1, 10.2],
                'NOM': [10.0, 10.0],
                '+TOL': [0.5, 0.5],
                '-TOL': [-0.5, -0.5],
            }
        )

        uncached = build_measurement_write_bundle('Diameter - X', header_group, 0)
        cache = {}
        cached = build_measurement_write_bundle_cached('Diameter - X', header_group, 0, cache=cache)

        self.assertEqual(cached['static_rows'], uncached['static_rows'])
        self.assertEqual(cached['header_plan']['stat_rows'], uncached['header_plan']['stat_rows'])
        self.assertEqual(cached['measurement_plan'], uncached['measurement_plan'])
        self.assertIn('usl_column', cached['measurement_plan'])
        self.assertIn('lsl_column', cached['measurement_plan'])

    def test_cached_header_plan_matches_uncached_output(self):
        header_group = pd.DataFrame(
            {
                'DATE': ['2024-01-01', '2024-01-02'],
                'SAMPLE_NUMBER': ['1', '2'],
                'MEAS': [10.1, 10.2],
                'NOM': [10.0, 10.0],
                '+TOL': [0.5, 0.5],
                '-TOL': [-0.5, -0.5],
            }
        )

        uncached = build_measurement_header_block_plan(header_group, 0)
        cache = {}
        cached_first = build_measurement_header_block_plan(header_group, 0, cache=cache)
        cached_second = build_measurement_header_block_plan(header_group, 0, cache=cache)

        self.assertEqual(cached_first, uncached)
        self.assertEqual(cached_second, uncached)
        self.assertEqual(len(cache.get('measurement_block_templates', {})), 1)

    def test_header_plan_exposes_cached_limit_formulas(self):
        header_group = pd.DataFrame(
            {
                'DATE': ['2024-01-01', '2024-01-02'],
                'SAMPLE_NUMBER': ['1', '2'],
                'MEAS': [10.1, 10.2],
                'NOM': [10.0, 10.0],
                '+TOL': [0.5, 0.5],
                '-TOL': [-0.5, -0.5],
            }
        )

        plan = build_measurement_header_block_plan(header_group, 0, cache={})
        self.assertEqual(plan['usl_formula'], '=(B1+B2)')
        self.assertEqual(plan['lsl_formula'], '=(B1+B3)')

    def test_create_measurement_formats_reuses_workbook_cache(self):
        from modules.export_sheet_writer import create_measurement_formats

        with tempfile.NamedTemporaryFile(suffix='.xlsx') as tmp:
            workbook = xlsxwriter.Workbook(tmp.name)
            try:
                first = create_measurement_formats(workbook)
                second = create_measurement_formats(workbook)
            finally:
                workbook.close()

        self.assertIs(first, second)


    def _write_owner_scoped_formats_workbook(self, path):
        from metroliza.exporting import export_sheet_writer as module
        from metroliza.exporting.export_backends import XlsxWorkbookAdapter

        with xlsxwriter.Workbook(str(path)) as workbook:
            adapter = XlsxWorkbookAdapter(workbook)
            adapter_ref = weakref.ref(adapter)
            formats = module.create_measurement_formats(adapter)
            owns_formats = all(value.xf_format_indices is workbook.xf_format_indices
                               for value in formats.values())
            worksheet = workbook.add_worksheet('MEASUREMENTS')
            worksheet.set_column(0, 3, 12, formats['default'])
            worksheet.write(0, 0, 'NOM')
            worksheet.write_number(1, 0, 0.25, formats['percent'])
            worksheet.write_string(0, 1, 'Date', formats['wrap'])
            worksheet.conditional_format(1, 0, 1, 0,
                                         {'type': 'cell', 'criteria': '>', 'value': 0,
                                          'format': formats['red']})
            del adapter
            self.assertIsNone(adapter_ref(), 'format cache must not retain ephemeral adapters')
        return formats, owns_formats

    def _assert_workbook_style_references_are_valid(self, path):
        from openpyxl import load_workbook

        namespace = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        with ZipFile(path) as package:
            styles = ElementTree.fromstring(package.read('xl/styles.xml'))
            count = len(styles.find('s:cellXfs', namespace))
            for member in package.namelist():
                if member.startswith('xl/worksheets/') and member.endswith('.xml'):
                    worksheet = ElementTree.fromstring(package.read(member))
                    for cell in worksheet.findall('.//s:c', namespace):
                        self.assertLess(int(cell.get('s', '0')), count, (member, cell.get('r')))
        workbook = load_workbook(path)
        try:
            self.assertEqual(workbook['MEASUREMENTS']['A1'].value, 'NOM')
        finally:
            workbook.close()

    def test_measurement_formats_do_not_alias_for_nonoverlapping_adapter_ids(self):
        from pathlib import Path
        from metroliza.exporting import export_sheet_writer as module

        # Force only the legacy cache key; no dependence on CPython allocation or explicit GC.
        with tempfile.TemporaryDirectory() as directory, \
                patch.object(module, 'id', lambda _owner: 41, create=True), \
                patch.object(module, '_MEASUREMENT_FORMATS_CACHE', {}, create=True):
            first, first_owned = self._write_owner_scoped_formats_workbook(Path(directory) / 'first.xlsx')
            second, second_owned = self._write_owner_scoped_formats_workbook(Path(directory) / 'second.xlsx')
            self.assertTrue(first_owned and second_owned)
            self.assertIsNot(first, second)
            self.assertIsNot(first['default'], second['default'])

    def test_measurement_styles_remain_valid_across_nonoverlapping_adapter_ids(self):
        from pathlib import Path
        from metroliza.exporting import export_sheet_writer as module

        with tempfile.TemporaryDirectory() as directory, \
                patch.object(module, 'id', lambda _owner: 41, create=True), \
                patch.object(module, '_MEASUREMENT_FORMATS_CACHE', {}, create=True):
            paths = [Path(directory) / name for name in ('first.xlsx', 'second.xlsx')]
            for path in paths:
                self._write_owner_scoped_formats_workbook(path)
            for path in paths:
                self._assert_workbook_style_references_are_valid(path)

    def test_measurement_formats_reuse_raw_owner_across_recreated_adapters(self):
        from pathlib import Path
        from metroliza.exporting import export_sheet_writer as module
        from metroliza.exporting.export_backends import XlsxWorkbookAdapter

        with tempfile.TemporaryDirectory() as directory, \
                xlsxwriter.Workbook(str(Path(directory) / 'owner.xlsx')) as workbook:
            first = module.create_measurement_formats(workbook)
            format_count = len(workbook.formats)
            for _index in range(3):
                adapter = XlsxWorkbookAdapter(workbook)
                self.assertIs(module.create_measurement_formats(adapter), first)
                del adapter
            self.assertIs(module.create_measurement_formats(workbook), first)
            self.assertEqual(len(workbook.formats), format_count)

    def test_measurement_formats_fall_back_without_attributes_hashing_or_weakrefs(self):
        from metroliza.exporting import export_sheet_writer as module

        calls = []

        class RestrictedWorkbook:
            __slots__ = ()
            __hash__ = None

            def add_format(self, properties):
                calls.append(dict(properties))
                return object()

        workbook = RestrictedWorkbook()
        first = module.create_measurement_formats(workbook)
        second = module.create_measurement_formats(workbook)
        self.assertIsNot(first, second)
        self.assertEqual(set(first), {'default', 'border', 'wrap', 'percent', 'red'})
        self.assertEqual(len(calls), 10)

    def test_measurement_format_cache_does_not_retain_released_owner_formats(self):
        from metroliza.exporting import export_sheet_writer as module

        class FormatToken:
            pass

        class WorkbookOwner:
            def add_format(self, _properties):
                return FormatToken()

        def create_references():
            owner = WorkbookOwner()
            formats = module.create_measurement_formats(owner)
            return weakref.ref(owner), weakref.ref(formats['default'])

        with patch.object(module, '_MEASUREMENT_FORMATS_CACHE', {}, create=True):
            owner_ref, format_ref = create_references()
            self.assertIsNone(owner_ref())
            self.assertIsNone(format_ref())


if __name__ == '__main__':
    unittest.main()
