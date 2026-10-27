"""Full-data numerical and XLSX equivalence for the #1128 export fast paths."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO
import json
import re
import warnings
import zipfile
import zlib

from matplotlib import cbook
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
import pytest

from metroliza.charts.full_data_violin import full_data_violin
from metroliza.exporting.xlsx_writer_policy import pandas_xlsxwriter_engine_kwargs
from metroliza.tabular.tabular_xlsx_writer import write_tabular_dataframe


@pytest.mark.parametrize('values', [
    np.arange(3.0), np.arange(150.0),
    np.random.default_rng(1128).normal(size=10001),
    np.array([-1e6, *np.zeros(200), 1e6]), np.ones(12),
    np.array([4.0]), np.array([np.nan, 1.0, 2.0, np.inf, 4.0]),
])
def test_violin_retains_full_density_vertices_and_statistics(values):
    """Every coordinate and summary statistic matches Matplotlib's full data KDE."""
    from metroliza.charts.full_data_violin import _univariate_density

    def reference(data, coords):
        from matplotlib.mlab import GaussianKDE
        if np.all(data[0] == data):
            return (data[0] == coords).astype(float)
        return GaussianKDE(data).evaluate(coords)

    expected = cbook.violin_stats([values], reference)
    actual = cbook.violin_stats([values], _univariate_density)
    assert expected[0].keys() == actual[0].keys()
    for key in expected[0]:
        np.testing.assert_array_equal(actual[0][key], expected[0][key], err_msg=key)


def test_violin_png_is_identical_to_matplotlib_full_data_render():
    values = [np.random.default_rng(1128).normal(size=211), np.ones(40)]
    images = []
    for renderer in (lambda ax, data, **kw: ax.violinplot(data, **kw), full_data_violin):
        fig, ax = plt.subplots()
        try:
            renderer(ax, values, positions=[1, 2], showmedians=True, showextrema=False)
            image = BytesIO()
            fig.savefig(image, format='png')
            images.append(image.getvalue())
        finally:
            plt.close(fig)
    assert images[0] == images[1]


def _frame_cases():
    yield pd.DataFrame({
        '=header': ['=1+1', 'https://example.invalid', '@SUM(A1)', None, ''],
        'float': [np.nan, np.inf, -np.inf, -0.0, 1.23456789012345],
        'integer': pd.Series([1, None, -3, 2**55, 0], dtype='Int64'),
        'boolean': pd.Series([True, False, None, True, False], dtype='boolean'),
        'datetime': pd.to_datetime(['2026-01-01', None, '2026-01-02', None, '2026-01-03']),
    })
    yield pd.DataFrame({'mixed': [date(2026, 1, 1), datetime(2026, 1, 1), timedelta(days=1), {'a': 1}, ['b']]})
    yield pd.DataFrame({'category': pd.Categorical(['a', None, 'b']), 'text': pd.Series(['a', None, ''], dtype='string')})
    yield pd.DataFrame(columns=['a', 'b'])
    yield pd.DataFrame()
    yield pd.DataFrame([[1, 'a']], columns=pd.MultiIndex.from_tuples([('a', 'x'), ('b', 'y')]))


def _workbook(frame, path, writer_function):
    with pd.ExcelWriter(path, engine='xlsxwriter', engine_kwargs=pandas_xlsxwriter_engine_kwargs()) as writer:
        writer_function(frame, writer, 'Data')


def _read_cells(path, sheet_name='Data'):
    book = openpyxl.load_workbook(path)
    try:
        return [[(cell.value, cell.data_type, cell.number_format, cell.style_id)
                 for cell in row] for row in book[sheet_name]]
    finally:
        book.close()


def _assert_workbooks_equal(expected, actual, *, sheet_names=('Data',)):
    with zipfile.ZipFile(expected) as before, zipfile.ZipFile(actual) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            expected_bytes, actual_bytes = before.read(name), after.read(name)
            # ZipFile.read also verifies each original member's stored CRC.
            if name == 'docProps/core.xml':
                expected_bytes = _without_core_timestamps(expected_bytes)
                actual_bytes = _without_core_timestamps(actual_bytes)
            else:
                assert before.getinfo(name).CRC == after.getinfo(name).CRC, name
            assert zlib.crc32(expected_bytes) == zlib.crc32(actual_bytes), name
            assert expected_bytes == actual_bytes, name
    for sheet_name in sheet_names:
        assert _read_cells(expected, sheet_name) == _read_cells(actual, sheet_name)


def _without_core_timestamps(payload):
    return re.sub(
        rb'(<dcterms:(created|modified) xsi:type="dcterms:W3CDTF">)'
        rb'\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z(</dcterms:\2>)',
        rb'\1TIMESTAMP\3', payload,
    )


def _rewrite_workbook_member(source, destination, member, old, new):
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(destination, 'w') as after:
        assert member in before.namelist()
        for info in before.infolist():
            payload = before.read(info.filename)
            if info.filename == member:
                assert payload.count(old) == 1, (member, old)
                payload = payload.replace(old, new, 1)
            after.writestr(info, payload)


@pytest.mark.parametrize(('member', 'old', 'new'), [
    ('xl/sharedStrings.xml', b'COLUMN_ORDER_A', b'COLUMN_ORDER_X'),
    ('xl/worksheets/sheet1.xml', b'<row r="2"', b'<row r="9"'),
    ('xl/worksheets/sheet1.xml', b'<v>12345</v>', b'<v>54321</v>'),
])
def test_workbook_comparison_rejects_shared_strings_rows_and_data(
    member, old, new, tmp_path,
):
    expected, changed = tmp_path / 'expected.xlsx', tmp_path / 'changed.xlsx'
    frame = pd.DataFrame({'text': ['COLUMN_ORDER_A'], 'number': [12345]})
    _workbook(frame, expected, write_tabular_dataframe)
    _rewrite_workbook_member(expected, changed, member, old, new)
    with pytest.raises(AssertionError, match=re.escape(member)):
        _assert_workbooks_equal(expected, changed)


def test_workbook_comparison_limits_timestamp_exception_to_core_fields(tmp_path):
    expected, changed = tmp_path / 'expected.xlsx', tmp_path / 'changed.xlsx'
    text = '2026-01-02T03:04:05Z'
    _workbook(pd.DataFrame({'timestamp': [text]}), expected, write_tabular_dataframe)
    _rewrite_workbook_member(
        expected, changed, 'xl/sharedStrings.xml', text.encode(),
        b'2026-01-02T03:04:06Z',
    )
    with pytest.raises(AssertionError, match='xl/sharedStrings.xml'):
        _assert_workbooks_equal(expected, changed)
    core = (
        b'<dcterms:created xsi:type="dcterms:W3CDTF">2026-01-02T03:04:05Z'
        b'</dcterms:created><dc:title>2026-01-02T03:04:05Z</dc:title>'
    )
    changed_core = core.replace(b'03:04:05Z</dcterms:', b'03:04:06Z</dcterms:')
    assert _without_core_timestamps(core) == _without_core_timestamps(changed_core)
    assert _without_core_timestamps(core) != _without_core_timestamps(
        core.replace(b'03:04:05Z</dc:title>', b'03:04:06Z</dc:title>')
    )


@pytest.mark.parametrize('rows', [0, 1, 18, 4097])
def test_tabular_writer_retains_column_major_shared_string_order(rows, tmp_path):
    frame = pd.DataFrame({
        'first': ['A_FIRST' if row % 2 == 0 else 'A_SECOND' for row in range(rows)],
        'second': ['B_FIRST' if row % 2 == 0 else 'A_FIRST' for row in range(rows)],
        'number': [float(row) for row in range(rows)],
    })
    expected, actual = tmp_path / 'expected.xlsx', tmp_path / 'actual.xlsx'
    _workbook(frame, expected, lambda df, writer, sheet: df.to_excel(writer, sheet_name=sheet, index=False))
    _workbook(frame, actual, write_tabular_dataframe)
    _assert_workbooks_equal(expected, actual)
    with zipfile.ZipFile(actual) as package:
        strings = package.read('xl/sharedStrings.xml')
    ordered = [b'first', b'second', b'number']
    if rows:
        ordered.append(b'A_FIRST')
    if rows > 1:
        ordered.append(b'A_SECOND')
    if rows:
        ordered.append(b'B_FIRST')
    assert [strings.index(b'<t>' + text + b'</t>') for text in ordered] == sorted(
        strings.index(b'<t>' + text + b'</t>') for text in ordered
    )


@pytest.mark.parametrize('guard', ['environment', 'internals', 'version', 'worksheet-api'])
def test_tabular_writer_preserves_workbook_when_bulk_guard_declines(guard, tmp_path, monkeypatch):
    from metroliza.exporting import xlsx_bulk_sheet

    if guard == 'environment':
        monkeypatch.setenv(xlsx_bulk_sheet.BULK_WRITE_ENV, '0')
    elif guard == 'internals':
        monkeypatch.setattr(xlsx_bulk_sheet, 'INTERNALS_AVAILABLE', False)
    elif guard == 'version':
        monkeypatch.setattr(xlsx_bulk_sheet.xlsxwriter, '__version__', '3.3.0')
    else:
        monkeypatch.setattr(xlsx_bulk_sheet, '_REQUIRED_WORKSHEET_ATTRIBUTES', ('missing_bulk_api',))
    assert not xlsx_bulk_sheet.bulk_write_enabled()
    frame = pd.DataFrame({
        'text': ['=literal', 'https://example.invalid', None],
        'number': [1.5, np.inf, np.nan],
        'boolean': pd.Series([True, None, False], dtype='boolean'),
    })
    expected, actual = tmp_path / 'expected.xlsx', tmp_path / 'actual.xlsx'
    _workbook(frame, expected, lambda df, writer, sheet: df.to_excel(writer, sheet_name=sheet, index=False))
    _workbook(frame, actual, write_tabular_dataframe)
    _assert_workbooks_equal(expected, actual)


def test_tabular_writer_activates_bulk_columns_without_per_cell_dispatch(tmp_path, monkeypatch):
    from xlsxwriter.worksheet import Worksheet

    from metroliza.exporting.xlsx_bulk_sheet import BulkRowsWorksheet, bulk_write_enabled

    monkeypatch.setenv('METROLIZA_EXPORT_BULK_WRITE', '1')
    if not bulk_write_enabled():
        pytest.skip('Installed XlsxWriter bulk API is unsupported')
    calls = []
    original_write = Worksheet._write

    def record_write(self, row, col, *args):
        calls.append((row, col))
        return original_write(self, row, col, *args)

    monkeypatch.setattr(Worksheet, '_write', record_write)
    frame = pd.DataFrame({'text': ['A', 'B', 'A'], 'number': [1.25, 2.5, -0.0]})
    with pd.ExcelWriter(tmp_path / 'activated.xlsx', engine='xlsxwriter',
                        engine_kwargs=pandas_xlsxwriter_engine_kwargs()) as writer:
        write_tabular_dataframe(frame, writer, 'Data')
        assert type(writer.sheets['Data']) is BulkRowsWorksheet
    assert calls == [(0, 0), (0, 1)]


@pytest.mark.parametrize('policy', ['existing-sheet', 'custom-sheet', 'constant-memory'])
def test_tabular_writer_preserves_existing_worksheet_policies(policy, tmp_path):
    from xlsxwriter.worksheet import Worksheet

    class CustomWorksheet(Worksheet):
        pass

    options = pandas_xlsxwriter_engine_kwargs()
    if policy == 'constant-memory':
        options['options']['constant_memory'] = True
    frame = pd.DataFrame({'text': ['A', 'B']})
    with pd.ExcelWriter(tmp_path / 'policy.xlsx', engine='xlsxwriter',
                        engine_kwargs=options) as writer:
        if policy == 'custom-sheet':
            writer.book.worksheet_class = CustomWorksheet
        elif policy == 'existing-sheet':
            existing = writer.book.add_worksheet('Data')
        write_tabular_dataframe(frame, writer, 'Data')
        expected_type = CustomWorksheet if policy == 'custom-sheet' else Worksheet
        assert type(writer.sheets['Data']) is expected_type
        if policy == 'existing-sheet':
            assert writer.sheets['Data'] is existing


@pytest.mark.parametrize('instance_override', [False, True])
def test_tabular_writer_preserves_class_worksheet_policy(instance_override, tmp_path, monkeypatch):
    from xlsxwriter import Workbook
    from xlsxwriter.worksheet import Worksheet

    class CustomWorksheet(Worksheet):
        pass

    monkeypatch.setattr(Workbook, 'worksheet_class', CustomWorksheet)
    frame = pd.DataFrame({'text': ['A', 'B'], 'number': [1.25, 2.5]})
    for enabled in ('0', '1'):
        monkeypatch.setenv('METROLIZA_EXPORT_BULK_WRITE', enabled)
        with pd.ExcelWriter(tmp_path / f'bulk-{enabled}.xlsx', engine='xlsxwriter',
                            engine_kwargs=pandas_xlsxwriter_engine_kwargs()) as writer:
            if instance_override:
                writer.book.worksheet_class = Worksheet
            write_tabular_dataframe(frame, writer, 'Data')
            sheet = writer.sheets['Data']
            if instance_override:
                assert isinstance(sheet, Worksheet) and not isinstance(sheet, CustomWorksheet)
            else:
                assert type(sheet) is CustomWorksheet
    _assert_workbooks_equal(tmp_path / 'bulk-0.xlsx', tmp_path / 'bulk-1.xlsx')


def test_tabular_writer_preserves_dataframe_subclass_header_callback(tmp_path):
    from xlsxwriter.worksheet import Worksheet

    observed_sheets = []

    class CustomFrame(pd.DataFrame):
        @property
        def _constructor(self):
            return CustomFrame

        def to_excel(self, writer, **kwargs):
            observed_sheets.append(tuple(writer.sheets))
            return super().to_excel(writer, **kwargs)

    with pd.ExcelWriter(tmp_path / 'subclass.xlsx', engine='xlsxwriter',
                        engine_kwargs=pandas_xlsxwriter_engine_kwargs()) as writer:
        write_tabular_dataframe(CustomFrame({'text': ['A']}), writer, 'Data')
        assert type(writer.sheets['Data']) is Worksheet
    assert observed_sheets == [()]


def _metadata_observer_workbook(path, case):
    events = []

    def event(operation):
        events.append(operation)
        warnings.warn(operation, UserWarning)

    class ObservedDict(dict):
        def __contains__(self, key):
            event('contains')
            return super().__contains__(key)

        def get(self, key, default=None):
            event('get')
            return super().get(key, default)

        def __getitem__(self, key):
            event('getitem')
            return super().__getitem__(key)

        def __iter__(self):
            event('iter')
            return super().__iter__()

    class ObservedStr(str):
        def __hash__(self):
            event('hash')
            return super().__hash__()

        def __eq__(self, other):
            event('eq')
            return super().__eq__(other)

    class ObservedFalse:
        def __bool__(self):
            event('bool')
            return False

    frame = pd.DataFrame({'text': ['A', 'B'], 'number': [1.25, 2.5]})
    with pd.ExcelWriter(path, engine='xlsxwriter',
                        engine_kwargs=pandas_xlsxwriter_engine_kwargs()) as writer:
        book = writer.book
        name = '' if case == 'empty' else 'Data'
        if case == 'mapping':
            book.sheetnames = ObservedDict(book.sheetnames)
        elif case == 'writer-metadata':
            writer.__dict__ = ObservedDict(writer.__dict__)
        elif case == 'book-metadata':
            book.__dict__ = ObservedDict(book.__dict__)
        elif case in {'writer-key', 'book-key'}:
            owner, key = (writer, '_book') if case == 'writer-key' else (book, 'sheetnames')
            owner.__dict__ = {ObservedStr(item) if item == key else item: value
                              for item, value in owner.__dict__.items()}
        elif case == 'sheet-key':
            sheet = book.add_worksheet('Existing')
            book.sheetnames = {ObservedStr('Existing'): sheet}
        elif case == 'name-subclass':
            name = ObservedStr(name)
        elif case == 'flag-subclass':
            book.constant_memory = ObservedFalse()
        events.clear()
        with warnings.catch_warnings(record=True) as captured:
            warnings.simplefilter('always')
            try:
                write_tabular_dataframe(frame, writer, name)
                error = None
            except Exception as exc:
                error = [type(exc).__name__, str(exc)]
            immediate = [sheet.name for sheet in book.worksheets()]
            immediate_types = [type(sheet).__name__ for sheet in book.worksheets()]
            if case == 'empty':
                write_tabular_dataframe(frame, writer, 'Valid')
            final_names = [str(sheet.name) for sheet in book.worksheets()]
            final_types = [type(sheet).__name__ for sheet in book.worksheets()]
        trace = {'events': list(events), 'warnings': [str(item.message) for item in captured],
                 'error': error, 'immediate': immediate, 'final_names': final_names,
                 'sheet_types': immediate_types}
        path.with_suffix('.worksheet-types.json').write_text(
            json.dumps({'immediate': immediate_types, 'final': final_types}), encoding='utf-8',
        )
    return trace


@pytest.mark.parametrize('case', [
    'empty', 'mapping', 'writer-metadata', 'book-metadata', 'writer-key',
    'book-key', 'sheet-key', 'name-subclass', 'flag-subclass',
])
def test_tabular_writer_preserves_metadata_observers(case, tmp_path, monkeypatch):
    traces = []
    for enabled in ('0', '1'):
        monkeypatch.setenv('METROLIZA_EXPORT_BULK_WRITE', enabled)
        traces.append(_metadata_observer_workbook(tmp_path / f'bulk-{enabled}.xlsx', case))
    (tmp_path / 'observer.json').write_text(json.dumps(traces, indent=2), encoding='utf-8')
    assert traces[0] == traces[1]
    if case == 'empty':
        assert traces[0]['error'] == ['KeyError', "''"]
        assert traces[0]['immediate'] == ['Sheet1']
        assert traces[0]['final_names'] == ['Sheet1', 'Valid']
    else:
        assert traces[0]['error'] is None
        assert all(kind == 'Worksheet' for kind in traces[0]['sheet_types'])
    _assert_workbooks_equal(tmp_path / 'bulk-0.xlsx', tmp_path / 'bulk-1.xlsx',
                            sheet_names=traces[0]['final_names'])


def _attrs_observer_workbook(path, mode, sheet_name, with_columns):
    with warnings.catch_warnings(record=True) as captured:
        warnings.simplefilter('always')
        with pd.ExcelWriter(path, engine='xlsxwriter',
                            engine_kwargs=pandas_xlsxwriter_engine_kwargs()) as writer:
            events = []

            class Attr:
                def __deepcopy__(self, memo):
                    state = list(writer.sheets)
                    events.append(state)
                    warnings.warn('attrs deepcopy sheets=' + repr(state), UserWarning)
                    if mode == 'raising':
                        raise RuntimeError('attrs deepcopy failed')
                    return self

            frame = pd.DataFrame({'number': [1, 2]}) if with_columns else pd.DataFrame(index=range(2))
            frame.attrs['probe'] = Attr()
            assert type(frame) is pd.DataFrame
            try:
                write_tabular_dataframe(frame, writer, sheet_name)
                error = None
            except Exception as exc:
                error = [type(exc).__name__, str(exc)]
            immediate = list(writer.sheets)
            write_tabular_dataframe(pd.DataFrame({'plain': [1, 2]}), writer, 'Valid')
            final_names = list(writer.sheets)
    return {'events': events, 'warnings': [str(item.message) for item in captured],
            'error': error, 'immediate': immediate, 'final_names': final_names}


@pytest.mark.parametrize(('mode', 'sheet_name', 'with_columns'), [
    pytest.param('raising', 'Data', False, id='zero-valid-raising'),
    pytest.param('observing', 'Data', False, id='zero-valid-observing'),
    pytest.param('raising', 'Invalid/Name', False, id='zero-invalid-raising'),
    pytest.param('observing', 'Invalid/Name', False, id='zero-invalid-observing'),
    pytest.param('raising', 'Data', True, id='normal-valid-raising'),
    pytest.param('observing', 'Data', True, id='normal-valid-observing'),
    pytest.param('raising', 'Invalid/Name', True, id='normal-invalid-raising'),
    pytest.param('observing', 'Invalid/Name', True, id='normal-invalid-observing'),
])
def test_tabular_writer_preserves_attrs_deepcopy_order(
    mode, sheet_name, with_columns, tmp_path, monkeypatch,
):
    traces = []
    for enabled in ('0', '1'):
        monkeypatch.setenv('METROLIZA_EXPORT_BULK_WRITE', enabled)
        traces.append(_attrs_observer_workbook(
            tmp_path / f'bulk-{enabled}.xlsx', mode, sheet_name, with_columns,
        ))
    (tmp_path / 'attrs-observer.json').write_text(json.dumps(traces, indent=2), encoding='utf-8')
    assert traces[0] == traces[1]
    if mode == 'raising':
        assert traces[0]['error'] == ['RuntimeError', 'attrs deepcopy failed']
        assert traces[0]['immediate'] == []
    elif sheet_name == 'Data':
        assert traces[0]['error'] is None
    else:
        assert traces[0]['error'][0] == 'InvalidWorksheetName'
    if not with_columns:
        assert traces[0]['events'] == [[]]
    _assert_workbooks_equal(tmp_path / 'bulk-0.xlsx', tmp_path / 'bulk-1.xlsx',
                            sheet_names=traces[0]['final_names'])


def test_tabular_writer_keeps_size_rejection_before_sheet_creation(tmp_path, monkeypatch):
    frame = pd.DataFrame(index=pd.RangeIndex(1_048_577))
    output = tmp_path / 'oversized.xlsx'
    with pd.ExcelWriter(output, engine='xlsxwriter',
                        engine_kwargs=pandas_xlsxwriter_engine_kwargs()) as writer:
        with pytest.raises(ValueError, match='sheet is too large'):
            write_tabular_dataframe(frame, writer, 'Data')
        assert not writer.sheets


def test_tabular_writer_preserves_accepted_empty_width_row_boundary(tmp_path):
    frame = pd.DataFrame(index=pd.RangeIndex(1_048_576))
    expected, actual = tmp_path / 'expected.xlsx', tmp_path / 'actual.xlsx'
    _workbook(frame, expected, lambda df, writer, sheet: df.to_excel(writer, sheet_name=sheet, index=False))
    _workbook(frame, actual, write_tabular_dataframe)
    _assert_workbooks_equal(expected, actual)


def test_analytics_workbook_all_data_sheets_match_disabled_bulk(tmp_path, monkeypatch):
    from metroliza.industrial.industrial_analytics_service import ProductionMetricCandidate
    from metroliza.tabular.tabular_analytics_service import export_tabular_analytics_workbook

    frame = pd.DataFrame({
        'source_row_number': [1, 2, 3],
        'reference': ['A', 'B', 'A'],
        'length': [1.25, 2.5, np.nan],
        'width': [4.0, 5.0, 6.0],
    })
    candidates = tuple(
        ProductionMetricCandidate(field, field.title(), 'dynamic', 3, 3, 1.0)
        for field in ('length', 'width')
    )
    results = []
    for enabled in ('0', '1'):
        monkeypatch.setenv('METROLIZA_EXPORT_BULK_WRITE', enabled)
        results.append(export_tabular_analytics_workbook(
            dataframe=frame, metric_candidates=candidates,
            output_file=tmp_path / f'bulk-{enabled}.xlsx', separate_parameter_sheets=True,
        ))
    assert results[0].sheet_names == results[1].sheet_names
    assert results[0].parameter_sheet_count == results[1].parameter_sheet_count == 2
    assert {'Table Data', 'Metrics', 'Diagnostics', 'Length', 'Width'} == set(results[0].sheet_names)
    _assert_workbooks_equal(results[0].output_file, results[1].output_file,
                            sheet_names=results[0].sheet_names)


@pytest.mark.parametrize('frame', list(_frame_cases())[:-1])
def test_bulk_writer_matches_pandas_cells_types_formats_and_xml(frame, tmp_path):
    expected, actual = tmp_path / 'expected.xlsx', tmp_path / 'actual.xlsx'
    _workbook(frame, expected, lambda df, writer, sheet: df.to_excel(writer, sheet_name=sheet, index=False))
    _workbook(frame, actual, write_tabular_dataframe)
    _assert_workbooks_equal(expected, actual)


@pytest.mark.parametrize('dtype', ['object', 'string'])
@pytest.mark.parametrize('text', ['x' * 32767, 'x' * 32768, '😀' * 32768],
                         ids=['at-limit', 'over-limit', 'unicode-over-limit'])
def test_bulk_writer_preserves_cells_after_oversized_text(dtype, text, tmp_path):
    frame = pd.DataFrame({'text': pd.Series(
        ['first', text, 'AFTER_LONG_TEXT', None, 'last'], dtype=dtype,
    )})
    expected, actual = tmp_path / 'expected.xlsx', tmp_path / 'actual.xlsx'
    _workbook(frame, expected, lambda df, writer, sheet: df.to_excel(writer, sheet_name=sheet, index=False))
    cells = _read_cells(expected)
    assert cells[2][0][0] == text[:32767]
    assert cells[3][0][0] == 'AFTER_LONG_TEXT'
    assert cells[5][0][0] == 'last'
    _workbook(frame, actual, write_tabular_dataframe)
    _assert_workbooks_equal(expected, actual)


@pytest.mark.parametrize('code', [-1, -2])
def test_bulk_writer_propagates_column_write_errors(code, tmp_path, monkeypatch):
    from xlsxwriter.worksheet import Worksheet

    monkeypatch.setenv('METROLIZA_EXPORT_BULK_WRITE', '0')
    monkeypatch.setattr(Worksheet, 'write_column', lambda *args: code)
    with pytest.raises(ValueError, match=f'XLSX column 0 write failed with code {code}'):
        _workbook(pd.DataFrame({'text': ['first', 'last']}), tmp_path / 'error.xlsx',
                  write_tabular_dataframe)


@pytest.mark.parametrize('missing', [None, pd.NA, np.nan, pd.NaT],
                         ids=['none', 'pd-na', 'nan', 'nat'])
def test_bulk_writer_preserves_object_numpy_boolean_cell_types(missing, tmp_path):
    frame = pd.DataFrame({'boolean': pd.Series(
        [np.bool_(True), missing, np.bool_(False), True, False], dtype=object,
    )})
    expected, actual = tmp_path / 'expected.xlsx', tmp_path / 'actual.xlsx'
    _workbook(frame, expected, lambda df, writer, sheet: df.to_excel(writer, sheet_name=sheet, index=False))
    cells = _read_cells(expected)
    assert [cells[row][0][:2] for row in (1, 3, 4, 5)] == [
        (True, 'b'), (False, 'b'), (True, 'b'), (False, 'b'),
    ]
    assert cells[2][0][0] is None
    _workbook(frame, actual, write_tabular_dataframe)
    _assert_workbooks_equal(expected, actual)


def test_bulk_writer_preserves_pandas_rejection_of_multiindex_columns(tmp_path):
    frame = list(_frame_cases())[-1]
    for writer_function in (write_tabular_dataframe, lambda df, writer, sheet: df.to_excel(writer, sheet_name=sheet, index=False)):
        with pytest.raises(NotImplementedError):
            _workbook(frame, tmp_path / 'multi.xlsx', writer_function)


def test_parallel_histogram_tables_match_serial_full_data_statistics():
    from metroliza.charts import hexafe_plotstats_adapter as adapter

    rng = np.random.default_rng(1128)
    groups = [(label, np.r_[rng.normal(center, 0.2, 2000), -100.0, 100.0])
              for label, center in [('B', 1.0), ('A', 2.0), ('C', 3.0)]]
    with adapter.histogram_stats_request():
        expected = [adapter.build_histogram_stats_table(values, title=label, lsl=0.5, usl=3.5)
                    for label, values in groups]
    with adapter.histogram_stats_request():
        actual = adapter.build_histogram_stats_tables(groups, lsl=0.5, usl=3.5)
        assert actual == expected
        assert [table.title for table in actual] == ['B', 'A', 'C']
        assert [dict(table.rows)['Samples'] for table in actual] == ['2002'] * 3
        # Workbook reuse after parallel dashboard computation must hit this cache.
        assert [adapter.build_histogram_stats_table(values, title=label, lsl=0.5, usl=3.5)
                for label, values in groups] == expected


def test_parallel_histogram_tables_share_bounded_cache_and_join_workers(monkeypatch):
    from threading import Barrier, enumerate as threads
    from metroliza.charts import hexafe_plotstats_adapter as adapter

    barrier = Barrier(3, timeout=10)
    calls = []

    def compute(values, **kwargs):
        calls.append(values[0])
        barrier.wait()
        return (('Samples', str(len(values))), ('Mean', str(float(np.mean(values)))))

    monkeypatch.setattr(adapter, '_histogram_table_rows_from_plotstats', compute)
    monkeypatch.setattr(adapter, '_HISTOGRAM_TABLE_CACHE_MAX_ENTRIES', 2)
    groups = [(str(index), np.arange(2000.0) + index) for index in range(3)]
    with adapter.histogram_stats_request():
        tables = adapter.build_histogram_stats_tables(groups)
        cache = adapter._histogram_table_cache.get()
        assert [table.title for table in tables] == ['0', '1', '2']
        assert len(cache.rows) == 2
        assert cache.payload_bytes <= adapter._HISTOGRAM_TABLE_CACHE_MAX_BYTES
        assert len(calls) == 3
    assert not cache.rows and cache.payload_bytes == 0
    assert not any(thread.name.startswith('histogram-stats') for thread in threads())


def test_parallel_histogram_failure_joins_workers_and_releases_request(monkeypatch):
    from threading import enumerate as threads
    from metroliza.charts import hexafe_plotstats_adapter as adapter

    def fail(*args, **kwargs):
        raise ValueError('synthetic fit failure')

    monkeypatch.setattr(adapter, '_histogram_table_rows_from_plotstats', fail)
    with pytest.raises(ValueError, match='synthetic fit failure'):
        with adapter.histogram_stats_request():
            cache = adapter._histogram_table_cache.get()
            adapter.build_histogram_stats_tables([('a', np.arange(3000.0)), ('b', np.arange(3000.0))])
    assert not cache.rows and cache.payload_bytes == 0
    assert adapter._histogram_table_cache.get() is None
    assert not any(thread.name.startswith('histogram-stats') for thread in threads())


def test_histogram_tables_fall_back_to_serial_if_threads_cannot_start(monkeypatch):
    from metroliza.charts import hexafe_plotstats_adapter as adapter

    class UnavailablePool:
        def __init__(self, **kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def submit(self, *args):
            raise RuntimeError('synthetic thread exhaustion')

    monkeypatch.setattr(adapter, 'ThreadPoolExecutor', UnavailablePool)
    monkeypatch.setattr(adapter, '_histogram_table_rows_from_plotstats',
                        lambda values, **kw: (('Samples', str(len(values))),))
    tables = adapter.build_histogram_stats_tables([('b', np.arange(3000.0)), ('a', np.arange(3000.0))])
    assert [table.title for table in tables] == ['b', 'a']
    assert [dict(table.rows)['Samples'] for table in tables] == ['3000', '3000']


def test_bulk_writer_retains_non_dataframe_diagnostics_payload(tmp_path):
    from metroliza.industrial.industrial_analytics_helpers import DiagnosticsSheetPayload

    payload = DiagnosticsSheetPayload(({'code': 'synthetic', 'message': '=literal', 'context': {'a': 1}},))
    expected, actual = tmp_path / 'expected.xlsx', tmp_path / 'actual.xlsx'
    _workbook(payload, expected, lambda data, writer, sheet: data.to_excel(writer, sheet_name=sheet, index=False))
    _workbook(payload, actual, write_tabular_dataframe)
    assert _read_cells(expected) == _read_cells(actual)


def test_histogram_pool_keeps_at_most_three_jobs_pending():
    from metroliza.charts.hexafe_plotstats_adapter import _bounded_histogram_tables

    class Future:
        def __init__(self, pool, job):
            self.pool, self.job = pool, job

        def result(self):
            self.pool.pending -= 1
            fn, *args = self.job
            return fn(*args)

    class Pool:
        pending = 0
        maximum = 0

        def submit(self, *job):
            self.pending += 1
            self.maximum = max(self.pending, self.maximum)
            assert self.pending <= 3
            return Future(self, job)

    pool = Pool()
    groups = list(range(100))
    assert _bounded_histogram_tables(pool, lambda group: group, groups) == groups
    assert pool.maximum == 3
    assert pool.pending == 0


def test_histogram_worker_runtime_error_is_not_mistaken_for_thread_exhaustion(monkeypatch):
    from metroliza.charts import hexafe_plotstats_adapter as adapter

    calls = []

    def fail(values, **kwargs):
        calls.append(1)
        raise RuntimeError('synthetic computation failure')

    monkeypatch.setattr(adapter, '_histogram_table_rows_from_plotstats', fail)
    with pytest.raises(RuntimeError, match='synthetic computation failure'):
        adapter.build_histogram_stats_tables([('a', np.arange(3000.0)), ('b', np.arange(3000.0))])
    assert len(calls) == 2
