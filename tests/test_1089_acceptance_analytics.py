"""Synthetic Reports-persistence to export acceptance checks for #1089."""

from __future__ import annotations

from contextlib import closing
import sqlite3
import zipfile

import pytest
from openpyxl import load_workbook
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QApplication, QDialog, QInputDialog

from metroliza.analytics.group_analysis_service import build_group_analysis_payload
from metroliza.exporting.contracts import AppPaths, ExportOptions, ExportRequest
from metroliza.exporting.export_outcomes import ExportRunStatus
from metroliza.reports.report_repository import ReportRepository
from metroliza.reports.report_query_service import build_measurement_filter_query
from metroliza.reports.report_schema import ensure_report_schema
from metroliza.ui.data_grouping import DataGrouping
from metroliza.exporting.export_data_thread import ExportDataThread


_APP = None


def _persist_report(repository, root, *, number, reference, header="FEATURE_1", ax="X", meas=10.1):
    source = root / f"public-synthetic-{number}.pdf"
    source.write_bytes(f"public-synthetic-report-{number}".encode())
    repository.persist_parsed_report(
        source_path=source,
        parser_id="cmm",
        parser_version="1.0",
        template_family="cmm_pdf_header_box",
        template_variant="cmm_pdf_header_box_serial_variant",
        parse_status="parsed_with_warnings",
        metadata={
            "reference": reference,
            "reference_raw": reference,
            "report_date": f"2024-01-0{number}",
            "part_name": f"Part {number}",
            "revision": "A",
            "sample_number": str(number),
            "sample_number_kind": "explicit_sample_number",
            "stats_count_raw": str(number),
            "stats_count_int": number,
            "operator_name": "Operator A",
            "comment": None,
        },
        measurements=[{
            "page_number": 1,
            "row_order": 1,
            "header": header,
            "section_name": header,
            "feature_label": header,
            "characteristic_name": "LOC",
            "characteristic_family": "LOC",
            "description": "Synthetic feature",
            "ax": ax,
            "nominal": 10.0,
            "tol_plus": 0.5,
            "tol_minus": -0.5,
            "bonus": 0.0,
            "meas": meas,
            "dev": meas - 10.0,
            "outtol": 0.0,
            "is_nok": False,
            "status_code": "ok",
        }],
        candidates=[],
        warnings=[],
        metadata_version="report_metadata_v1",
        metadata_profile_id="cmm_pdf_header_box",
        metadata_profile_version="1",
        page_count=1,
        measurement_count=1,
        has_nok=False,
        nok_count=0,
        metadata_confidence=0.5,
        identity_hash=f"synthetic-identity-{number}",
    )


def _fixture_db(tmp_path):
    db_path = tmp_path / "public-synthetic.sqlite"
    ensure_report_schema(str(db_path))
    repository = ReportRepository(str(db_path))
    for number, reference in ((1, None), (2, ""), (3, "REF-3")):
        _persist_report(
            repository,
            tmp_path,
            number=number,
            reference=reference,
            header="FEATURE_B" if number == 2 else "FEATURE_A",
            ax="Y" if number == 2 else "X",
            meas=10.0 + number / 10,
        )
    return db_path


def _workbook_parts(path):
    with zipfile.ZipFile(path) as workbook:
        return sorted(workbook.namelist())


def test_missing_reference_reports_keep_analytical_workbook_and_dashboard(tmp_path):
    db_path = _fixture_db(tmp_path)
    with closing(sqlite3.connect(db_path)) as connection:
        assert connection.execute("SELECT COUNT(*) FROM report_measurements").fetchone()[0] == 3
        assert connection.execute(
            "SELECT COUNT(*) FROM report_metadata WHERE reference IS NULL OR TRIM(reference) = ''"
        ).fetchone()[0] == 2

    workbook_path = tmp_path / "public-synthetic.xlsx"
    request = ExportRequest(
        paths=AppPaths(db_file=str(db_path), excel_file=str(workbook_path)),
        options=ExportOptions(generate_summary_sheet=True, generate_html_dashboard=True),
    )
    thread = ExportDataThread(request)
    thread.run()

    assert workbook_path.exists()
    parts = _workbook_parts(workbook_path)
    assert len([part for part in parts if part.startswith("xl/charts/chart")]) == 3
    workbook = load_workbook(workbook_path)
    assert workbook["MEASUREMENTS"].max_row == 4
    raw_rows = list(workbook['MEASUREMENTS'].values)
    reference_index = raw_rows[0].index('REFERENCE')
    raw_references = [row[reference_index] for row in raw_rows[1:]]
    assert raw_references.count('REF-3') == 1
    assert sum(value in (None, '') for value in raw_references) == 2
    assert sum(bool(workbook[sheet]._charts) for sheet in workbook.sheetnames) == 3
    assert "Sheet" not in workbook.sheetnames
    assert any(name.startswith("Report 1") for name in workbook.sheetnames)
    assert any(name.startswith("Report 2") for name in workbook.sheetnames)
    assert thread.completion_metadata["html_dashboard_path"]
    assert thread.completion_metadata["html_dashboard_section_count"] == 3
    assert thread.completion_metadata["html_dashboard_chart_count"] > 0


def test_grouping_can_select_and_apply_two_named_groups_to_export(tmp_path, monkeypatch):
    global _APP
    db_path = _fixture_db(tmp_path)
    _persist_report(
        ReportRepository(str(db_path)), tmp_path, number=4,
        reference='REF-3', header='FEATURE_A', ax='X', meas=10.4,
    )
    _APP = QApplication.instance() or _APP or QApplication([])
    app = _APP

    class Parent(QDialog):
        def __init__(self):
            super().__init__()
            self.df_for_grouping = None
            self.grouping_applied = False

        def set_df_for_grouping(self, frame):
            self.df_for_grouping = frame

        def set_grouping_applied(self, applied):
            self.grouping_applied = applied

    parent = Parent()
    dialog = DataGrouping(parent=parent, db_file=str(db_path))
    try:
        assert len(dialog.df) == 4
        assert dialog.reference_list.count() == 3
        for index in range(dialog.reference_list.count()):
            dialog.reference_list.setCurrentRow(index)
            app.processEvents()
            assert dialog.part_list.count() >= 1
        for report_number, group_name in ((1, "Alpha"), (3, "Alpha"), (4, "Beta")):
            target_key = dialog.df.loc[dialog.df['REPORT_ID'] == report_number, 'GROUP_KEY'].iloc[0]
            for index in range(dialog.reference_list.count()):
                dialog.reference_list.setCurrentRow(index)
                app.processEvents()
                for part_index in range(dialog.part_list.count()):
                    if dialog.part_list.item(part_index).data(Qt.ItemDataRole.UserRole) == target_key:
                        dialog.part_list.setCurrentRow(part_index)
                        break
                else:
                    continue
                break
            else:
                raise AssertionError(f"Report {report_number} unavailable for grouping")
            monkeypatch.setattr(QInputDialog, "getText", lambda *_a, _name=group_name, **_k: (_name, True))
            dialog.create_group()
        dialog.use_grouping()
        assert parent.grouping_applied is True
        assert parent.df_for_grouping is not None
        assigned = parent.df_for_grouping.set_index('REPORT_ID')['GROUP'].to_dict()
        assert assigned[1] == 'Alpha'
        assert assigned[3] == 'Alpha'
        assert assigned[4] == 'Beta'

        workbook_path = tmp_path / 'grouped.xlsx'
        request = ExportRequest(
            paths=AppPaths(db_file=str(db_path), excel_file=str(workbook_path)),
            options=ExportOptions(generate_html_dashboard=True, group_analysis_level='standard'),
            grouping_df=parent.df_for_grouping,
        )
        thread = ExportDataThread(request)
        captured = {}
        original_write = thread._write_html_dashboard_if_requested

        def capture_group_payload():
            original_write()
            captured['group_analysis'] = thread._html_group_analysis_payload

        thread._write_html_dashboard_if_requested = capture_group_payload
        thread.run()
        assert workbook_path.exists()
        workbook = load_workbook(workbook_path)
        assert workbook['MEASUREMENTS'].max_row == 5
        assert sum(bool(workbook[sheet]._charts) for sheet in workbook.sheetnames) == 3
        assert not any(sheet.endswith('_summary') for sheet in workbook.sheetnames)
        assert thread.completion_metadata['html_dashboard_section_count'] == 3
        assert captured['group_analysis']['diagnostics']['group_count'] >= 2
        assert captured['group_analysis']['metric_rows'], captured['group_analysis']
        assert thread.export_run_result.status is ExportRunStatus.COMPLETE_WITH_OMISSIONS
        assert 'without a reference' in thread.completion_metadata['group_analysis_warnings'][0]
        dashboard = (tmp_path / 'grouped_dashboard.html').read_text(encoding='utf-8')
        assert 'Alpha' in dashboard and 'Beta' in dashboard

        complete_path = tmp_path / 'complete-groups.xlsx'
        complete_thread = ExportDataThread(
            ExportRequest(
                paths=AppPaths(db_file=str(db_path), excel_file=str(complete_path)),
                options=ExportOptions(generate_html_dashboard=True, group_analysis_level='standard'),
                filter_query=build_measurement_filter_query(reference_values=('REF-3',)),
                grouping_df=parent.df_for_grouping,
            )
        )
        complete_thread.run()
        assert complete_thread.export_run_result.status is ExportRunStatus.COMPLETE
        assert complete_thread.completion_metadata['html_dashboard_section_count'] == 1
        assert load_workbook(complete_path)['MEASUREMENTS'].max_row == 3
    finally:
        dialog.close()
        parent.close()


@pytest.mark.parametrize('missing_reference', [None, '', '   '])
def test_group_analysis_never_compares_known_with_missing_reference(missing_reference):
    payload = build_group_analysis_payload([
        {'REFERENCE': 'REF-A', 'HEADER - AX': 'FEATURE_A - X', 'GROUP': 'Alpha', 'MEAS': 10.0},
        {'REFERENCE': missing_reference, 'HEADER - AX': 'FEATURE_A - X', 'GROUP': 'Beta', 'MEAS': 10.1},
    ])
    assert payload['metric_rows'] == []
    assert payload['status'] == 'skipped'
    assert payload['diagnostics']['reference_count'] == 1


def test_group_analysis_with_only_missing_references_has_explicit_skip_reason():
    payload = build_group_analysis_payload([
        {'REFERENCE': None, 'HEADER - AX': 'FEATURE_A - X', 'GROUP': 'Alpha', 'MEAS': 10.0},
        {'REFERENCE': '', 'HEADER - AX': 'FEATURE_A - X', 'GROUP': 'Beta', 'MEAS': 10.1},
    ])
    assert payload['status'] == 'skipped'
    assert payload['metric_rows'] == []
    assert payload['skip_reason']['code'] == 'missing_reference_metadata'


def test_grouped_export_keeps_missing_report_charts_without_false_comparison(tmp_path):
    db_path = tmp_path / 'one-known-one-missing.sqlite'
    ensure_report_schema(str(db_path))
    repository = ReportRepository(str(db_path))
    _persist_report(repository, tmp_path, number=1, reference=None)
    _persist_report(repository, tmp_path, number=2, reference='REF-A', meas=10.2)
    workbook_path = tmp_path / 'one-known-one-missing.xlsx'
    thread = ExportDataThread(
        ExportRequest(
            paths=AppPaths(db_file=str(db_path), excel_file=str(workbook_path)),
            options=ExportOptions(generate_html_dashboard=True, group_analysis_level='standard'),
            grouping_df=(
                {'REPORT_ID': 1, 'GROUP': 'Beta'},
                {'REPORT_ID': 2, 'GROUP': 'Alpha'},
            ),
        )
    )
    captured = {}
    original_write = thread._write_html_dashboard_if_requested

    def capture_group_payload():
        original_write()
        captured['group_analysis'] = thread._html_group_analysis_payload

    thread._write_html_dashboard_if_requested = capture_group_payload
    thread.run()

    assert workbook_path.exists()
    workbook = load_workbook(workbook_path)
    assert workbook['MEASUREMENTS'].max_row == 3
    assert sum(bool(workbook[sheet]._charts) for sheet in workbook.sheetnames) == 2
    assert thread.completion_metadata['html_dashboard_section_count'] == 2
    assert captured['group_analysis']['metric_rows'] == []
    assert thread.export_run_result.status is ExportRunStatus.COMPLETE_WITH_OMISSIONS
    assert any('without a reference' in warning for warning in thread.completion_metadata['group_analysis_warnings'])


def test_filters_constrain_measurements_and_analytical_partitions_together(tmp_path):
    db_path = _fixture_db(tmp_path)
    cases = (
        ('all', {}, 3),
        ('reference', {'reference_values': ('REF-3',)}, 1),
        ('header_axis', {'header_values': ('FEATURE_B',), 'ax_values': ('Y',)}, 1),
        ('expression', {'expression_text': 'Measurement>=10.2'}, 2),
        ('metadata', {'part_name_values': ('Part 2',)}, 1),
    )
    for name, fields, expected in cases:
        workbook_path = tmp_path / f'filtered-{name}.xlsx'
        thread = ExportDataThread(
            ExportRequest(
                paths=AppPaths(db_file=str(db_path), excel_file=str(workbook_path)),
                options=ExportOptions(),
                filter_query=build_measurement_filter_query(**fields),
            )
        )
        thread.run()
        workbook = load_workbook(workbook_path)
        assert workbook['MEASUREMENTS'].max_row - 1 == expected, name
        assert sum(bool(workbook[sheet]._charts) for sheet in workbook.sheetnames) == expected, name


def test_html_only_keeps_missing_reference_sections(tmp_path):
    db_path = _fixture_db(tmp_path)
    html_path = tmp_path / 'dashboard-only.html'
    thread = ExportDataThread(
        ExportRequest(
            paths=AppPaths(db_file=str(db_path), html_dashboard_file=str(html_path)),
            options=ExportOptions(
                preset='html_dashboard_only',
                export_target='html_dashboard',
                backend_target='html',
                generate_summary_sheet=True,
                generate_html_dashboard=True,
            ),
        )
    )
    thread.run()
    assert html_path.exists()
    assert thread.export_run_result.status is ExportRunStatus.COMPLETE
    assert thread.completion_metadata['html_dashboard_section_count'] == 3
    assert thread.completion_metadata['html_dashboard_chart_count'] > 0
    html = html_path.read_text(encoding='utf-8')
    assert 'Report 1 (no reference)' in html
    assert 'Report 2 (no reference)' in html
    assert not list(tmp_path.glob('*.xlsx'))


@pytest.mark.parametrize('dashboard_requested', [False, True])
def test_grouped_export_with_no_comparable_metric_reports_omission(tmp_path, dashboard_requested):
    db_path = _fixture_db(tmp_path)
    workbook_path = tmp_path / 'uncomparable-groups.xlsx'
    thread = ExportDataThread(
        ExportRequest(
            paths=AppPaths(db_file=str(db_path), excel_file=str(workbook_path)),
            options=ExportOptions(
                generate_html_dashboard=dashboard_requested,
                group_analysis_level='standard',
            ),
            grouping_df=(
                {'REPORT_ID': 1, 'GROUP': 'Alpha'},
                {'REPORT_ID': 2, 'GROUP': 'Default'},
                {'REPORT_ID': 3, 'GROUP': 'Beta'},
            ),
        )
    )
    thread.run()
    assert workbook_path.exists()
    if dashboard_requested:
        assert thread.completion_metadata['html_dashboard_section_count'] == 3
    assert thread.export_run_result.status is ExportRunStatus.COMPLETE_WITH_OMISSIONS
    assert 'no metric with measurements from two groups' in (
        thread.completion_metadata['group_analysis_warnings'][0]
    )


def test_nonempty_scope_without_analytical_keys_fails_before_workbook_publication(tmp_path):
    db_path = tmp_path / 'incomplete.sqlite'
    ensure_report_schema(str(db_path))
    _persist_report(
        ReportRepository(str(db_path)), tmp_path,
        number=1, reference=None, header=None,
    )
    workbook_path = tmp_path / 'incomplete.xlsx'
    thread = ExportDataThread(
        ExportRequest(
            paths=AppPaths(db_file=str(db_path), excel_file=str(workbook_path)),
            options=ExportOptions(),
        )
    )
    with pytest.raises(ValueError, match='no usable reference/report partition with HEADER and AX'):
        thread.get_export_backend().run(thread)
    assert not workbook_path.exists()
